import json
import re
from pathlib import Path

from openpyxl import load_workbook

from app.batch_fast_prelisting import _output_filename
from app.template_red_field_scanner import scan_template_red_fields


ROUND_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = (ROUND_DIR / "../outputs").resolve()
MANIFEST = json.loads((ROUND_DIR / "manifest.json").read_text(encoding="utf-8"))


F = {
    "model": "model_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "style": "style[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "department": "department[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "gender": "target_gender[marketplace_id=ATVPDKIKX0DER]#1.value",
    "age": "age_range_description[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "fabric": "fabric_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "care": "care_instructions[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "import": "import_designation[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "size": "size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "unit": "unit_count[marketplace_id=ATVPDKIKX0DER]#1.value",
    "unit_type": "unit_count[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US].value",
    "included": "included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "special": "special_feature[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "warranty": "warranty_description[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
}


ROW_VALUES = {
    22: {
        F["model"]: "Lace Brief Panties", "bottom_style[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Briefs",
        F["style"]: "Lace Brief", F["department"]: "Womens", F["gender"]: "Female", F["age"]: "Adult",
        "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_system": "US", "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_class": "Alpha",
        F["fabric"]: "85% Polyamide, 15% Spandex", F["care"]: "Hand Wash Only",
        "rise[marketplace_id=ATVPDKIKX0DER]#1.style[language_tag=en_US]#1.value": "Low",
        "pattern[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Solid", F["import"]: "Imported",
    },
    23: {
        F["model"]: "Cotton Boxer Briefs", "bottom_style[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Boxer Briefs",
        F["style"]: "Boxer Brief", F["department"]: "Womens", F["gender"]: "Female", F["age"]: "Adult",
        "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_system": "US", "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_class": "Alpha",
        F["fabric"]: "100% Cotton", F["care"]: "Machine Wash",
        "rise[marketplace_id=ATVPDKIKX0DER]#1.style[language_tag=en_US]#1.value": "Mid",
        "pattern[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Solid", F["import"]: "Imported",
    },
    37: {
        F["model"]: "Full Face Cycling Mask", F["style"]: "Full Face", F["department"]: "Unisex",
        F["gender"]: "Unisex", F["age"]: "Adult", "headwear_size[marketplace_id=ATVPDKIKX0DER]#1.size_system": "US",
        "headwear_size[marketplace_id=ATVPDKIKX0DER]#1.size_class": "Alpha", F["fabric"]: "Polyester",
        F["care"]: "Hand Wash Only", "seasons[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "All", F["import"]: "Imported",
    },
    38: {
        F["model"]: "Thermal Neck Gaiter", F["style"]: "Neck Gaiter", F["department"]: "Unisex",
        F["gender"]: "Unisex", F["age"]: "Adult", "headwear_size[marketplace_id=ATVPDKIKX0DER]#1.size_system": "US",
        "headwear_size[marketplace_id=ATVPDKIKX0DER]#1.size_class": "Alpha", F["fabric"]: "Polyester",
        F["care"]: "Hand Wash Only", "seasons[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Winter", F["import"]: "Imported",
    },
    41: {F["model"]: "Adhesive Toilet Seat Cover", F["style"]: "Adhesive", F["included"]: "Toilet Seat Cover",
         "item_length_width[marketplace_id=ATVPDKIKX0DER]#1.length.value": 3.94,
         "item_length_width[marketplace_id=ATVPDKIKX0DER]#1.length.unit": "Inches",
         "item_length_width[marketplace_id=ATVPDKIKX0DER]#1.width.value": 3.94,
         "item_length_width[marketplace_id=ATVPDKIKX0DER]#1.width.unit": "Inches",
         "number_of_boxes[marketplace_id=ATVPDKIKX0DER]#1.value": 1},
    42: {F["model"]: "Cotton Ankle Socks", F["style"]: "Ankle", F["department"]: "Unisex", F["gender"]: "Unisex", F["age"]: "Adult",
         "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_system": "US", "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_class": "Alpha",
         "item_weight[marketplace_id=ATVPDKIKX0DER]#1.value": 0.22},
    44: {F["model"]: "Cooling Arm Sleeves", F["style"]: "Sports Sleeve", F["department"]: "Unisex", F["gender"]: "Unisex", F["age"]: "Adult",
         F["fabric"]: "Polyester", F["size"]: "One Size", F["care"]: "Hand Wash Only",
         "closure[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value": "Pull On", F["import"]: "Imported"},
    47: {F["model"]: "Divided Bento Box", F["special"]: "Divided", "capacity[marketplace_id=ATVPDKIKX0DER]#1.value": 2,
         "capacity[marketplace_id=ATVPDKIKX0DER]#1.unit": "Cups", F["unit"]: 2, F["unit_type"]: "Count",
         F["included"]: "2 Bento Boxes with Lids", "item_volume[marketplace_id=ATVPDKIKX0DER]#1.value": 2,
         "item_volume[marketplace_id=ATVPDKIKX0DER]#1.unit": "Cups"},
    49: {F["model"]: "Round Ice Cube Tray", F["size"]: "Standard", "item_shape[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Round",
         F["care"]: "Hand Wash Only", "is_dishwasher_safe[marketplace_id=ATVPDKIKX0DER]#1.value": "No", F["unit"]: 1,
         F["unit_type"]: "Count", F["included"]: "Ice Cube Tray, Lid and Storage Container"},
    50: {F["model"]: "Gradient Arm Sleeves", F["style"]: "Gradient", F["department"]: "Unisex", F["gender"]: "Unisex", F["age"]: "Adult",
         F["fabric"]: "Polyester", F["size"]: "One Size", F["care"]: "Hand Wash Only",
         "closure[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value": "Pull On", F["import"]: "Imported"},
    53: {F["model"]: "Precision Tweezer Set", F["unit"]: 4, F["unit_type"]: "Count"},
    71: {F["model"]: "Portable Sewing Kit", F["warranty"]: "No Warranty"},
    72: {F["model"]: "Cycling Athletic Shorts", F["style"]: "Cycling", F["department"]: "Men", F["gender"]: "Male", F["age"]: "Adult",
         "bottoms_size[marketplace_id=ATVPDKIKX0DER]#1.size_system": "US", "bottoms_size[marketplace_id=ATVPDKIKX0DER]#1.size_class": "Alpha",
         F["fabric"]: "100% Polyester", "item_length_description[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Above Knee",
         F["care"]: "Machine Wash", "rise[marketplace_id=ATVPDKIKX0DER]#1.style[language_tag=en_US]#1.value": "Mid Rise",
         "closure[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value": "Pull On", F["import"]: "Imported"},
    78: {F["model"]: "Magnetizer Demagnetizer", F["special"]: "Portable", F["unit"]: 1, F["unit_type"]: "Count",
         "head[marketplace_id=ATVPDKIKX0DER]#1.style[language_tag=en_US]#1.value": "Standard"},
    79: {F["model"]: "SAE Ball End Hex Key Set", "finish_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Polished",
         F["unit"]: 9, F["unit_type"]: "Count", F["included"]: "9 Hex Keys and Holder",
         "head[marketplace_id=ATVPDKIKX0DER]#1.style[language_tag=en_US]#1.value": "Ball End"},
    81: {F["model"]: "Precision Repair Kit", F["special"]: "Flexible Shaft", F["unit"]: 23, F["unit_type"]: "Count",
         "head[marketplace_id=ATVPDKIKX0DER]#1.style[language_tag=en_US]#1.value": "Multi-Bit"},
    82: {"target_audience_keyword[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Women",
         "generic_keyword[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "metal hair tie ponytail holder bow geometric accessory",
         "lifestyle[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Fashion", F["gender"]: "Female",
         "number_of_pieces[marketplace_id=ATVPDKIKX0DER]#1.value": 1, F["unit"]: 1, F["unit_type"]: "Count",
         "hair_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "All"},
    86: {F["model"]: "Summer Sport Socks", F["style"]: "Low Cut", F["department"]: "Men", F["gender"]: "Male", F["age"]: "Adult",
         "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_system": "US", "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_class": "Alpha",
         "item_weight[marketplace_id=ATVPDKIKX0DER]#1.value": 0.22},
    87: {F["model"]: "Fine Mesh Strainer", F["style"]: "Spoon Strainer", F["care"]: "Hand Wash Only",
         "is_dishwasher_safe[marketplace_id=ATVPDKIKX0DER]#1.value": "No", F["included"]: "Mesh Strainer"},
    88: {F["model"]: "Summer Boat Socks", F["style"]: "No Show", F["department"]: "Men", F["gender"]: "Male", F["age"]: "Adult",
         "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_system": "US", "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_class": "Alpha",
         "item_weight[marketplace_id=ATVPDKIKX0DER]#1.value": 0.22},
    91: {F["model"]: "Rug Corner Gripper", "item_shape[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Triangle",
         "item_length_width[marketplace_id=ATVPDKIKX0DER]#1.length.value": 3.15, "item_length_width[marketplace_id=ATVPDKIKX0DER]#1.length.unit": "Inches",
         "item_length_width[marketplace_id=ATVPDKIKX0DER]#1.width.value": 3.15, "item_length_width[marketplace_id=ATVPDKIKX0DER]#1.width.unit": "Inches"},
    92: {"measurement_accuracy[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "1 MM"},
    93: {F["model"]: "Exercise Resistance Band", F["special"]: "Lightweight", F["department"]: "Unisex", F["unit"]: 1, F["unit_type"]: "Count",
         F["included"]: "Exercise Band", "sport_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Exercise and Fitness",
         "recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Strength Training", F["import"]: "Imported", F["warranty"]: "No Warranty"},
    94: {"handle[marketplace_id=ATVPDKIKX0DER]#1.length#1.decimal_value": 9,
         "handle[marketplace_id=ATVPDKIKX0DER]#1.length#1.string_value": "9 Inches", F["included"]: "Leather Hole Punch Pliers"},
    95: {F["model"]: "Multi Function Bottle Opener", F["size"]: "Standard", F["care"]: "Hand Wash Only",
         "is_customizable[marketplace_id=ATVPDKIKX0DER]#1.value": "No", F["included"]: "3 Bottle Openers"},
    96: {F["special"]: "360 Degree Rotation", "compatible_phone_models[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Universal Smartphones",
         "compatible_devices[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Smartphones",
         "mounting_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Dashboard", F["included"]: "Phone Mount", F["warranty"]: "No Warranty"},
    97: {F["model"]: "Shoe Laundry Bag", F["special"]: "Reusable", "item_shape[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Rectangular",
         F["unit"]: 1, F["unit_type"]: "Count", F["included"]: "Shoe Wash Bag", "closure[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value": "Zipper"},
    101: {F["model"]: "Shoe Cleaning Bag", F["special"]: "Mesh", "item_shape[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Rectangular",
          F["unit"]: 1, F["unit_type"]: "Count", F["included"]: "Shoe Wash Bag", "closure[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value": "Zipper"},
    106: {F["model"]: "S Shape Snap Hook", "lock_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Wire Lock",
          F["included"]: "20 Snap Hooks", F["import"]: "Imported", F["warranty"]: "No Warranty", F["fabric"]: "Metal"},
    110: {"compatible_phone_models[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Universal Smartphones",
          "compatible_devices[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Smartphones",
          "item_thickness[marketplace_id=ATVPDKIKX0DER]#1.decimal_value": 0.04,
          "item_thickness[marketplace_id=ATVPDKIKX0DER]#1.string_value": "0.04 Inches",
          "number_of_boxes[marketplace_id=ATVPDKIKX0DER]#1.value": 1, F["warranty"]: "No Warranty"},
    115: {F["model"]: "Cherry Keychain", F["special"]: "Lightweight", F["department"]: "Women", F["gender"]: "Female", F["size"]: "Standard",
          "theme[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Cherry", "closure[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value": "Split Ring", F["import"]: "Imported"},
    116: {"jewelry_material_categorization[marketplace_id=ATVPDKIKX0DER]#1.value": "Fashion", F["size"]: "One Size",
          "gem_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Crystal", "stones[marketplace_id=ATVPDKIKX0DER]#1.id": 1,
          "stones[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US].value": "Crystal",
          "stones[marketplace_id=ATVPDKIKX0DER]#1.creation_method[language_tag=en_US].value": "Simulated",
          "stones[marketplace_id=ATVPDKIKX0DER]#1.treatment_method[language_tag=en_US].value": "Not Treated",
          "stones[marketplace_id=ATVPDKIKX0DER]#1.clarity[language_tag=en_US].value": "Included",
          "stones[marketplace_id=ATVPDKIKX0DER]#1.color[language_tag=en_US].value": "White",
          "stones[marketplace_id=ATVPDKIKX0DER]#1.cut[language_tag=en_US].value": "Good",
          "stones[marketplace_id=ATVPDKIKX0DER]#1.shape[language_tag=en_US].value": "Round",
          "clasp_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Lobster Claw",
          "item_thickness[marketplace_id=ATVPDKIKX0DER]#1.decimal_value": 0.2,
          "item_thickness[marketplace_id=ATVPDKIKX0DER]#1.string_value": "0.2 Inches",
          "item_length[marketplace_id=ATVPDKIKX0DER]#1.value": 7.5,
          "item_length[marketplace_id=ATVPDKIKX0DER]#1.unit": "Inches",
          "item_weight[marketplace_id=ATVPDKIKX0DER]#1.value": 0.22,
          "item_weight[marketplace_id=ATVPDKIKX0DER]#1.unit": "Pounds",
          "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.length.value": 3.94,
          "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.length.unit": "Inches",
          "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.width.value": 3.94,
          "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.width.unit": "Inches",
          "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.height.value": 0.79,
          "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.height.unit": "Inches",
          "cpsia_cautionary_statement[marketplace_id=ATVPDKIKX0DER]#1.value": "No Warning Applicable"},
    117: {F["model"]: "Mosquito Coil Holder", "generic_keyword[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "mosquito coil holder incense burner patio camping",
          "item_shape[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Round",
          "item_length_width_thickness[marketplace_id=ATVPDKIKX0DER]#1.length.value": 5.91,
          "item_length_width_thickness[marketplace_id=ATVPDKIKX0DER]#1.length.unit": "Inches",
          "item_length_width_thickness[marketplace_id=ATVPDKIKX0DER]#1.thickness.value": 1.18,
          "item_length_width_thickness[marketplace_id=ATVPDKIKX0DER]#1.thickness.unit": "Inches",
          "item_length_width_thickness[marketplace_id=ATVPDKIKX0DER]#1.width.value": 5.91,
          "item_length_width_thickness[marketplace_id=ATVPDKIKX0DER]#1.width.unit": "Inches",
          "required_product_compliance_certificate[marketplace_id=ATVPDKIKX0DER]#1.value": "Not Applicable"},
    118: {F["department"]: "Women", F["size"]: "Standard", "metal_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Alloy",
          "metals[marketplace_id=ATVPDKIKX0DER]#1.id": 1, "metals[marketplace_id=ATVPDKIKX0DER]#1.metal_type[language_tag=en_US].value": "Alloy",
          "metals[marketplace_id=ATVPDKIKX0DER]#1.metal_stamp[language_tag=en_US].value": "No Metal Stamp",
          "metals[marketplace_id=ATVPDKIKX0DER]#1.metal_weight.value": 10,
          "metals[marketplace_id=ATVPDKIKX0DER]#1.metal_weight.unit": "Grams",
          "gem_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "No Gemstone"},
    119: {F["model"]: "Capybara Keychain", F["special"]: "Lightweight", F["department"]: "Unisex", F["gender"]: "Unisex", F["size"]: "Standard",
          "theme[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Capybara", "closure[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value": "Split Ring", F["import"]: "Imported"},
    122: {F["model"]: "Blank Resin Keychain", F["special"]: "Customizable", F["department"]: "Unisex", F["gender"]: "Unisex", F["size"]: "Standard",
          "theme[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "DIY", "closure[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value": "Split Ring", F["import"]: "Imported"},
    124: {F["model"]: "Square Parchment Sheets", "number_of_pieces[marketplace_id=ATVPDKIKX0DER]#1.value": 50,
          F["care"]: "Single Use Only", "sheet_count[marketplace_id=ATVPDKIKX0DER]#1.value": 50},
}


def field_value(row_number, task, sku, field_name):
    if field_name == "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size":
        if sku.endswith("-S"): return "Small"
        if sku.endswith("-M") or "-M-" in sku: return "Medium"
        return "One Size"
    if field_name == "headwear_size[marketplace_id=ATVPDKIKX0DER]#1.size":
        return "One Size"
    if field_name == "bottoms_size[marketplace_id=ATVPDKIKX0DER]#1.size":
        if sku.endswith("-M"): return "Medium"
        if sku.endswith("-2XL"): return "XX-Large (xx_l)"
        return "Medium"
    if field_name == "set_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value":
        if "DoubleArc" in sku: return "Double Arc"
        if "Bow" in sku: return "Bow"
        if "BallSet" in sku: return "2 Piece Ball Set"
        if "MixedSet" in sku: return "3 Piece Mixed Set"
        if "RopeKnot" in sku: return "1 Piece Long Rope Knot"
    values = ROW_VALUES.get(row_number, {})
    if field_name in values:
        return values[field_name]
    if field_name == F["model"]: return task["product_name"]
    if field_name == F["unit"]: return task.get("set_count") or 1
    if field_name == F["unit_type"]: return "Count"
    if field_name == F["included"]: return task["product_name"]
    if field_name == F["special"]: return "Lightweight"
    if field_name == F["warranty"]: return "No Warranty"
    return None


def clean_listing_text(ws, fields, row_number):
    listing_fields = [
        "item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
        "product_description[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    ] + [f"bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#{i}.value" for i in range(1, 6)]
    for rr in range(7, ws.max_row + 1):
        for field_name in listing_fields:
            col = fields.get(field_name)
            if not col or not isinstance(ws.cell(rr, col).value, str):
                continue
            value = ws.cell(rr, col).value
            replacements = [
                (r"\bBreathable\b", "Lightweight"), (r"\bbreathable\b", "lightweight"),
                (r"\bTherapy\b", "Exercise"), (r"\btherapy\b", "exercise"),
                (r"\bMagnetic\b", "Magnet"), (r"\bmagnetic\b", "magnet"),
                (r"\bBin\b", "Container"), (r"\bbin\b", "container"),
                (r"30\s*[Cc][Mm]\b", "11.81 Inches"), (r"40\s*[Mm][Mm]\b", "40 MM"),
                (r"8\s*[xX×]\s*6", "8 by 6"), (r"\baverage size\b", "Average Size"),
            ]
            for pattern, replacement in replacements:
                value = re.sub(pattern, replacement, value)
            if field_name.startswith("item_name"):
                for word, replacement in [("socks", "Footwear"), ("shorts", "Sportswear"), ("tool", "Accessory")]:
                    count = 0
                    def sub(match):
                        nonlocal count
                        count += 1
                        return match.group(0) if count <= 2 else replacement
                    value = re.sub(rf"\b{word}\b", sub, value, flags=re.I)
            ws.cell(rr, col).value = value


ROW_ORDER = [22, 23, 37, 38, 41, 42, 44, 47, 48, 49, 50, 53, 71, 72, 78, 79, 81, 82, 86, 87, 88, 91, 92, 93, 94, 95, 96, 97, 101, 106, 110, 112, 115, 116, 117, 118, 119, 122, 124]


for row_number, task in zip(ROW_ORDER, MANIFEST["tasks"]):
    path = OUTPUT_DIR / _output_filename(task, task["name"], Path(task["template"]).suffix)
    wb = load_workbook(path, keep_vba=True)
    ws = wb["Template"]
    fields = {str(ws.cell(5, col).value).strip(): col for col in range(1, ws.max_column + 1) if ws.cell(5, col).value}
    clean_listing_text(ws, fields, row_number)
    wb.save(path)

    # Fill exactly the cells currently triggered by the template's conditional rules.
    for _ in range(4):
        red, _unresolved = scan_template_red_fields(path)
        if not red:
            break
        wb = load_workbook(path, keep_vba=True)
        ws = wb["Template"]
        fields = {str(ws.cell(5, col).value).strip(): col for col in range(1, ws.max_column + 1) if ws.cell(5, col).value}
        sku_col = fields["contribution_sku#1.value"]
        changed = 0
        for item in red:
            field_name = item["field_name"]
            value = field_value(row_number, task, item.get("sku", ""), field_name)
            col = fields.get(field_name)
            if col and value not in (None, ""):
                ws.cell(item["row"], col).value = value
                changed += 1
        if not changed:
            break
        wb.save(path)

    # The four tweezers are sold as one confirmed 4-piece bundle, not as a variation family.
    if row_number == 53:
        wb = load_workbook(path, keep_vba=True)
        ws = wb["Template"]
        fields = {str(ws.cell(5, col).value).strip(): col for col in range(1, ws.max_column + 1) if ws.cell(5, col).value}
        sku_col = fields["contribution_sku#1.value"]
        parentage = fields.get("parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value")
        parent_sku = fields.get("child_parent_sku_relationship[marketplace_id=ATVPDKIKX0DER]#1.parent_sku")
        theme = fields.get("variation_theme#1.name")
        for rr in range(7, ws.max_row + 1):
            sku = str(ws.cell(rr, sku_col).value or "")
            if sku == "TTCA-TweezerSet":
                for col in range(1, ws.max_column + 1): ws.cell(rr, col).value = None
            elif sku == "TTCA-TweezerSet-4pcs":
                if parentage: ws.cell(rr, parentage).value = None
                if parent_sku: ws.cell(rr, parent_sku).value = None
                if theme: ws.cell(rr, theme).value = None
        wb.save(path)

print("patched", len(MANIFEST["tasks"]), "outputs")
