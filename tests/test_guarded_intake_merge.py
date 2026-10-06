"""Synthetic-only checks for the opt-in guarded default-intake merge."""

import hashlib
import json
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app.guarded_intake import GuardedIntakeConflict, guarded_merge_intake
from app import workbench


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class GuardedIntakeMergeTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / "data" / "projects"
        self.project = self.root / "P-SYN"
        self.intake_dir = self.project / "03_产品详情页"
        self.intake_dir.mkdir(parents=True)
        self.draft = self.intake_dir / "synthetic.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "产品资料"
        sheet.append(["sku", "product_name", "cost", "list_price", "haul_price", "supplier_link"])
        sheet.append(["SKU-A", "Synthetic", "", "", "", ""])
        book.save(self.draft)
        book.close()
        self.status = self.project / "project_status.json"
        self.status.write_text(json.dumps({"status": "not_started",
            "latest_draft": "03_产品详情页/synthetic.xlsx", "sku_count": 1}), encoding="utf-8")

    def payload(self):
        return {"approved": True, "project_ref": "data/projects/P-SYN",
            "expected_status_sha256": sha(self.status), "expected_draft_sha256": sha(self.draft),
            "expected_skus": ["SKU-A"], "changes": [
                {"sku": "SKU-A", "field": "list_price", "value": 1.88},
                {"sku": "SKU-A", "field": "haul_price", "value": 1.88},
                {"sku": "SKU-A", "field": "cost", "value": 2.5},
                {"sku": "SKU-A", "field": "supplier_link", "value": "https://example.invalid/item"},
            ]}

    def apply(self, payload):
        return guarded_merge_intake(payload, projects_root=self.root)

    def test_fill_empty_and_recoverable_backup(self):
        request = self.payload()
        old_status, old_draft = sha(self.status), sha(self.draft)
        result = self.apply(request)
        self.assertTrue(result["ok"])
        self.assertEqual(result["changed_cells"], 4)
        self.assertEqual(sha(self.status), old_status)
        self.assertEqual(sha(Path(result["backup_file"])), old_draft)
        self.assertNotEqual(sha(self.draft), old_draft)
        book = load_workbook(self.draft, read_only=True, data_only=True)
        row = [cell.value for cell in book.active[2]]
        book.close()
        self.assertEqual(row[2:], [2.5, 1.88, 1.88, "https://example.invalid/item"])
        with self.draft.open("ab") as changed_source:
            changed_source.write(b"synthetic later edit")
        self.assertEqual(sha(Path(result["backup_file"])), old_draft)
        with self.assertRaises(GuardedIntakeConflict):
            self.apply(request)

    def test_rejects_stale_status_and_nonblank_value(self):
        request = self.payload()
        before = sha(self.draft)
        self.status.write_text(self.status.read_text(encoding="utf-8") + "\n", encoding="utf-8")
        with self.assertRaises(GuardedIntakeConflict):
            self.apply(request)
        self.assertEqual(sha(self.draft), before)
        request = self.payload()
        book = load_workbook(self.draft)
        book.active["D2"] = 1.74
        book.save(self.draft)
        book.close()
        request["expected_draft_sha256"] = sha(self.draft)
        with self.assertRaises(GuardedIntakeConflict):
            self.apply(request)

    def test_rejects_extra_sku_and_invalid_field(self):
        request = self.payload()
        request["expected_skus"] = ["SKU-A", "SKU-B"]
        with self.assertRaises(GuardedIntakeConflict):
            self.apply(request)
        request = self.payload()
        request["changes"][0]["field"] = "status"
        with self.assertRaises(ValueError):
            self.apply(request)

    def test_atomic_replace_failure_leaves_source_unchanged(self):
        request = self.payload()
        before = sha(self.draft)
        with patch("app.guarded_intake.os.replace", side_effect=OSError("synthetic replace failure")):
            with self.assertRaises(OSError):
                self.apply(request)
        self.assertEqual(sha(self.draft), before)
        self.assertFalse(list(self.intake_dir.glob(".guarded_intake_*.xlsx")))

    def test_http_route_uses_guard_and_returns_conflict(self):
        original = workbench.guarded_merge_intake
        workbench.guarded_merge_intake = lambda payload: self.apply(payload)
        server = ThreadingHTTPServer(("127.0.0.1", 0), workbench._handler())
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            request = self.payload()
            url = f"http://127.0.0.1:{server.server_port}/api/guarded-intake-merge"
            data = json.dumps(request).encode("utf-8")
            with urllib.request.urlopen(urllib.request.Request(url, data=data,
                    headers={"Content-Type": "application/json"})) as response:
                self.assertTrue(json.load(response)["ok"])
            with self.assertRaises(urllib.error.HTTPError) as error:
                urllib.request.urlopen(urllib.request.Request(url, data=data,
                    headers={"Content-Type": "application/json"}))
            self.assertEqual(error.exception.code, 409)
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)
            workbench.guarded_merge_intake = original


if __name__ == "__main__":
    unittest.main()
