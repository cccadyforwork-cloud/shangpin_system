"""Create the cable V2 by changing only the rejected item-length cells."""

from pathlib import Path
import re
import sys
import tempfile
import zipfile

from openpyxl import load_workbook


def main(source: Path, target: Path) -> None:
    if target.exists():
        raise FileExistsError(target)

    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=True)
    if workbook.sheetnames[4] != "Template":
        raise ValueError("Template sheet is not sheet5.xml")
    sheet = workbook["Template"]
    if sheet["EE4"].value != "Item Length" or sheet["EF4"].value != "Item Length Unit":
        raise ValueError("Item length columns have changed")

    replacements = {}
    for row in range(8, 14):
        sku = sheet[f"A{row}"].value
        meters = 1 if row % 2 == 0 else 2
        if not str(sku).endswith(f"-{meters}m"):
            raise ValueError(f"Unexpected SKU at row {row}: {sku}")
        if sheet[f"EE{row}"].value != meters or sheet[f"EF{row}"].value != "Meters":
            raise ValueError(f"Unexpected item length at row {row}")
        replacements[f"EE{row}"] = (str(meters), "39.37" if meters == 1 else "78.74")
        replacements[f"EF{row}"] = ("Meters", "Inches")

    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".xlsm", delete=False) as temp:
        temporary = Path(temp.name)
    try:
        with zipfile.ZipFile(source) as input_zip, zipfile.ZipFile(temporary, "w") as output_zip:
            for member in input_zip.infolist():
                data = input_zip.read(member.filename)
                if member.filename == "xl/worksheets/sheet5.xml":
                    xml = data.decode("utf-8")
                    for cell, (old, new) in replacements.items():
                        pattern = rf'(<c r="{cell}"[^>]*><(?:v|is><t)>){re.escape(old)}(</(?:v|t></is)></c>)'
                        xml, count = re.subn(pattern, lambda match: match.group(1) + new + match.group(2), xml)
                        if count != 1:
                            raise ValueError(f"Expected one replacement at {cell}, found {count}")
                    data = xml.encode("utf-8")
                output_zip.writestr(member, data)
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
