import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Border, Side
from openpyxl.workbook.defined_name import DefinedName

from app.template_red_field_scanner import scan_template_red_fields


class TemplateRedFieldScannerTest(unittest.TestCase):
    def test_reports_only_triggered_blank_child_cell(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Template"
        sheet.cell(4, 1).value = "SKU"
        sheet.cell(4, 2).value = "Product Type"
        sheet.cell(4, 3).value = "Model Name"
        sheet.cell(4, 4).value = "Parentage"
        sheet.cell(5, 1).value = "contribution_sku#1.value"
        sheet.cell(5, 2).value = "product_type#1.value"
        sheet.cell(5, 3).value = "model_name#1.value"
        sheet.cell(5, 4).value = "parentage_level#1.value"
        sheet.append([])
        sheet.cell(7, 1).value = "SKU-CHILD-BLANK"
        sheet.cell(7, 2).value = "TYPE"
        sheet.cell(7, 4).value = "Child"
        sheet.cell(8, 1).value = "SKU-CHILD-FILLED"
        sheet.cell(8, 2).value = "TYPE"
        sheet.cell(8, 3).value = "Filled"
        sheet.cell(8, 4).value = "Child"
        sheet.cell(9, 1).value = "SKU-PARENT"
        sheet.cell(9, 2).value = "TYPE"
        sheet.cell(9, 4).value = "Parent"

        workbook.defined_names.add(DefinedName("PTList0", attr_text='Template!$B1="TYPE"'))
        workbook.defined_names.add(DefinedName("form_rgPTList0field", attr_text='Template!$D1="Parent"'))
        workbook.defined_names.add(DefinedName("reqPTList0field", attr_text='NOT(LEN(Template!$C1)>0)'))

        grey = Side(style="thin", color="FFD4D4D4")
        green = Side(style="thin", color="FF008000")
        red = Side(style="thin", color="FFFF0000")
        sheet.conditional_formatting.add(
            "C7:C9",
            FormulaRule(formula=["IF(PTList0,form_rgPTList0field,0)"], stopIfTrue=True,
                        border=Border(left=grey, right=grey, top=grey, bottom=grey)),
        )
        sheet.conditional_formatting.add(
            "C7:C9",
            FormulaRule(formula=["IF(LEN(C7)>0,1,0)"], stopIfTrue=True,
                        border=Border(left=green, right=green, top=green, bottom=green)),
        )
        sheet.conditional_formatting.add(
            "C7:C9",
            FormulaRule(formula=["IF(PTList0,reqPTList0field,0)"], stopIfTrue=True,
                        border=Border(left=red, right=red, top=red, bottom=red)),
        )

        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "template.xlsx"
            workbook.save(path)
            findings, unresolved = scan_template_red_fields(path)

        self.assertEqual(unresolved, [])
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["coordinate"], "C7")
        self.assertEqual(findings[0]["sku"], "SKU-CHILD-BLANK")
        self.assertEqual(findings[0]["label"], "Model Name")


if __name__ == "__main__":
    unittest.main()
