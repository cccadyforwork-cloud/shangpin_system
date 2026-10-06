from openpyxl import Workbook

from app.template_validator import FIELD_NAMES, validate_template_file


def _make_template(path, apparel_size):
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"
    headers = [
        FIELD_NAMES["sku"],
        "product_type#1.value",
        FIELD_NAMES["condition"],
        FIELD_NAMES["parentage_level"],
        FIELD_NAMES["fabric_type"],
        FIELD_NAMES["target_gender"],
        FIELD_NAMES["age_range"],
        FIELD_NAMES["apparel_size_system"],
        FIELD_NAMES["apparel_size_class"],
        FIELD_NAMES["apparel_size"],
    ]
    for column, field_name in enumerate(headers, 1):
        ws.cell(5, column).value = field_name
    values = [
        "TEST-UNDERPANTS-M",
        "UNDERPANTS",
        "New",
        "Child",
        "100% Cotton",
        "Female",
        "Adult",
        "US",
        "Alpha",
        apparel_size,
    ]
    for column, value in enumerate(values, 1):
        ws.cell(7, column).value = value
    wb.save(path)


def test_underpants_rejects_abbreviated_apparel_size(tmp_path):
    path = tmp_path / "invalid.xlsx"
    _make_template(path, "M")

    findings, _ = validate_template_file(path)

    assert any(item["field"] == "Apparel Size Value" for item in findings)


def test_underpants_accepts_exact_apparel_size_enum(tmp_path):
    path = tmp_path / "valid.xlsx"
    _make_template(path, "Medium")

    findings, _ = validate_template_file(path)

    assert not any(item["field"] == "Apparel Size Value" for item in findings)
