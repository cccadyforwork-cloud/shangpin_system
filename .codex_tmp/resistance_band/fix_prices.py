from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook


SOURCE = Path("/Users/cc/Documents/GitHub/shangpin_system/outputs/01a09a0d-5e1d-7473-8611-f221c4c977bf/弹力带价格刷新V2.xlsx")
OUTPUT = Path("/Users/cc/Documents/GitHub/shangpin_system/outputs/01a09a0d-5e1d-7473-8611-f221c4c977bf/弹力带价格刷新V3.xlsx")

FIELD_LIST_PRICE = "list_price[marketplace_id=ATVPDKIKX0DER]#1.value"
FIELD_HAUL_PRICE = "purchasable_offer[marketplace_id=ATVPDKIKX0DER][audience=BZR]#1.our_price#1.schedule#1.value_with_tax"
FIELD_MIN_PRICE = "purchasable_offer[marketplace_id=ATVPDKIKX0DER][audience=BZR]#1.minimum_seller_allowed_price#1.schedule#1.value_with_tax"
FIELD_MAX_PRICE = "purchasable_offer[marketplace_id=ATVPDKIKX0DER][audience=BZR]#1.maximum_seller_allowed_price#1.schedule#1.value_with_tax"
FIELD_SKIP_OFFER = "skip_offer[marketplace_id=ATVPDKIKX0DER]#1.value"
FIELD_PARENTAGE = "parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value"


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
copyfile(SOURCE, OUTPUT)
workbook = load_workbook(OUTPUT, keep_links=True)
sheet = workbook["模板"]
fields = {
    str(sheet.cell(5, col).value).strip(): col
    for col in range(1, sheet.max_column + 1)
    if sheet.cell(5, col).value
}

required = [FIELD_LIST_PRICE, FIELD_HAUL_PRICE, FIELD_MIN_PRICE, FIELD_SKIP_OFFER, FIELD_PARENTAGE]
missing = [field for field in required if field not in fields]
if missing:
    raise ValueError(f"Missing required price fields: {missing}")

for row in range(7, 11):
    is_parent = str(sheet.cell(row, fields[FIELD_PARENTAGE]).value or "").strip().lower() == "parent"
    sheet.cell(row, fields[FIELD_SKIP_OFFER]).value = None
    sheet.cell(row, fields[FIELD_MIN_PRICE]).value = None
    if FIELD_MAX_PRICE in fields:
        sheet.cell(row, fields[FIELD_MAX_PRICE]).value = None
    if is_parent:
        sheet.cell(row, fields[FIELD_LIST_PRICE]).value = None
        sheet.cell(row, fields[FIELD_HAUL_PRICE]).value = None
    else:
        sheet.cell(row, fields[FIELD_LIST_PRICE]).value = 2.59
        sheet.cell(row, fields[FIELD_HAUL_PRICE]).value = 2.59

workbook.calculation.fullCalcOnLoad = True
workbook.calculation.forceFullCalc = True
workbook.calculation.calcMode = "auto"
workbook.save(OUTPUT)
print(OUTPUT)
