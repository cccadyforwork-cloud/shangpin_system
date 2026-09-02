import json
import html as html_lib
import re
import tempfile
import zipfile
from copy import deepcopy
from itertools import combinations
from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook

from .analyzer import analyze_project
from .paths import COMPETITOR_DIR, CONFIG_DIR, PRODUCT_DETAIL_DIR, TEMPLATE_SOURCE_DIR, safe_name
from .template_validator import validate_template_file
from .template_red_field_scanner import scan_template_red_fields
from .template_writer import fill_template
from .workbook_io import read_intake_rows, write_intake_workbook


LOGISTICS_TIERS_PATH = CONFIG_DIR / "logistics_tiers.json"
FAST_ROUTE = "Haul Generic Variation"
REFERENCE_SAFE_FIELDS = {
    "package_level[marketplace_id=ATVPDKIKX0DER]#1.value",
    "model_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "style[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "number_of_pieces[marketplace_id=ATVPDKIKX0DER]#1.value",
    "item_shape[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "pattern[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "item_form[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "unit_count[marketplace_id=ATVPDKIKX0DER]#1.value",
    "unit_count[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US].value",
    "recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value",
    "recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#3.value",
    "number_of_packs[marketplace_id=ATVPDKIKX0DER]#1.value",
    "product_tax_code#1.value",
}


def run_batch_fast_prelisting(manifest_path, output_dir=None):
    manifest_path = Path(manifest_path).expanduser().resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    tasks = manifest.get("tasks") or []
    if not tasks:
        raise ValueError("批量清单没有 tasks。")

    base_dir = manifest_path.parent
    output_dir = Path(output_dir).expanduser().resolve() if output_dir else _resolve_optional_path(
        manifest.get("output_dir"), base_dir
    )
    if output_dir is None:
        output_dir = base_dir / "批量快速上品输出"
    output_dir.mkdir(parents=True, exist_ok=True)

    results = []
    for index, task in enumerate(tasks, 1):
        try:
            result = _run_one_task(task, index, base_dir, output_dir)
        except Exception as exc:
            result = {
                "index": index,
                "name": str(task.get("name") or task.get("output_name") or f"任务{index}"),
                "status": "failed",
                "error_count": None,
                "output": None,
                "message": str(exc),
            }
        results.append(result)

    return {
        "manifest": manifest_path,
        "output_dir": output_dir,
        "task_count": len(tasks),
        "success_count": sum(1 for item in results if item["status"] == "ready_for_review"),
        "failed_count": sum(1 for item in results if item["status"] == "failed"),
        "needs_fix_count": sum(1 for item in results if item["status"] == "needs_manual_fix"),
        "needs_wps_count": sum(1 for item in results if item["status"] == "needs_wps"),
        "results": results,
    }


def _run_one_task(task, index, base_dir, output_dir):
    task = deepcopy(task)
    name = str(task.get("name") or task.get("output_name") or f"任务{index}").strip()
    template_path = _required_path(task, "template", base_dir)
    competitor_paths = _competitor_paths(task, base_dir)
    competitor_title = extract_competitor_title(competitor_paths)
    if competitor_title:
        task["_competitor_base_title"] = rewrite_competitor_title_tail(competitor_title)
    if not task.get("variants") and task.get("expand_competitor_variants", True):
        detected_variants = extract_competitor_variants(competitor_paths)
        if detected_variants:
            task["variants"] = detected_variants
            task.pop("child_sku", None)
    price = _resolve_task_price(task, competitor_paths)
    tier = resolve_logistics_tier(task.get("logistics_tier"), task.get("weight_grams"))
    copy_defaults = _reference_copy(task.get("copy_reference"), base_dir)
    reference_fields = _reference_safe_fields(task.get("copy_reference"), base_dir)
    reference_rows = _reference_rows(task.get("copy_reference"), base_dir)

    with tempfile.TemporaryDirectory(prefix="shangpin_fast_") as temp_name:
        project_dir = Path(temp_name) / safe_name(task.get("product_name") or name)
        template_dir = project_dir / TEMPLATE_SOURCE_DIR
        competitor_dir = project_dir / COMPETITOR_DIR
        detail_dir = project_dir / PRODUCT_DETAIL_DIR
        template_dir.mkdir(parents=True, exist_ok=True)
        competitor_dir.mkdir(parents=True, exist_ok=True)
        detail_dir.mkdir(parents=True, exist_ok=True)

        local_template = template_dir / template_path.name
        copyfile(template_path, local_template)
        for competitor_path in competitor_paths:
            copyfile(competitor_path, competitor_dir / competitor_path.name)

        draft_path, _report_path = analyze_project(
            project_dir,
            detail_dir / f"{safe_name(name)}_自动提炼草稿.xlsx",
            route_mode="variation",
            write_report=False,
        )
        rows = read_intake_rows(draft_path)
        rows = _apply_fast_overrides(rows, task, name, price, tier, local_template, copy_defaults)
        write_intake_workbook(draft_path, rows)

        output_name = _output_filename(task, name, template_path.suffix)
        filled_path = output_dir / output_name
        fill_template(
            project_dir,
            draft_path=draft_path,
            template_path=local_template,
            output_path=filled_path,
            write_report=False,
        )
        _apply_batch_overlay(filled_path, task, tier, reference_fields, reference_rows)
        findings, _ = validate_template_file(filled_path, write_report=False)
        red_fields, unresolved_red_rules = scan_template_red_fields(filled_path)
        _check_zip_structure(filled_path)

    error_count = len(findings)
    if error_count:
        status = "needs_manual_fix"
        message = findings[0]["message"]
    elif unresolved_red_rules:
        status = "needs_wps"
        message = f"有 {len(unresolved_red_rules)} 个条件格式公式暂不能计算，请用 WPS 复核。"
    elif red_fields:
        status = "needs_manual_fix"
        first = red_fields[0]
        message = f"检测到 {len(red_fields)} 个红框字段；首项：{first['sku']} / {first['label']}。"
    else:
        status = "ready_for_review"
        message = "项目自检与红框扫描均通过，可进入人工复核。"
    return {
        "index": index,
        "name": name,
        "status": status,
        "error_count": error_count,
        "red_field_count": len(red_fields),
        "unresolved_red_rule_count": len(unresolved_red_rules),
        "output": str(filled_path),
        "message": message,
    }


def extract_competitor_variants(paths):
    """Extract a complete one-dimensional color twister from saved Amazon HTML.

    The competitor ASINs are intentionally not returned because this route creates
    new Generic listings and does not attach to the competitor relationship.
    """
    variants = []
    seen = set()
    for path in paths:
        text = Path(path).read_text(encoding="utf-8", errors="ignore")
        key_match = re.search(r'name=["\']twisterDimKeys["\']\s+value=["\']([^"\']+)', text, re.I)
        if not key_match or key_match.group(1).strip() != "color_name":
            continue
        pattern = re.compile(
            r'<li\b[^>]*\bid=["\']color_name_\d+["\'][^>]*\btitle=["\']Click to select ([^"\']+)["\'][^>]*>',
            re.I,
        )
        for match in pattern.finditer(text):
            color = _normalize_competitor_color(html_lib.unescape(match.group(1)))
            key = color.casefold()
            if not color or key in seen:
                continue
            seen.add(key)
            variants.append({"color": color})
    return variants


def extract_competitor_price(paths):
    for path in paths:
        text = Path(path).read_text(encoding="utf-8", errors="ignore")
        match = re.search(
            r'class=["\'][^"\']*aok-offscreen[^"\']*["\']>\s*\$([0-9][0-9,]*(?:\.[0-9]{1,2})?)\s*</span>\s*'
            r'<span\b[^>]*class=["\'][^"\']*priceToPay[^"\']*["\']',
            text,
            re.I,
        )
        if match:
            return float(match.group(1).replace(",", ""))
    return None


def extract_competitor_title(paths):
    for path in paths:
        text = Path(path).read_text(encoding="utf-8", errors="ignore")
        match = re.search(
            r'<[^>]+\bid=["\']productTitle["\'][^>]*>(.*?)</[^>]+>',
            text,
            re.I | re.S,
        )
        if not match:
            match = re.search(r"<title[^>]*>(.*?)</title>", text, re.I | re.S)
        if not match:
            continue
        title = re.sub(r"<[^>]+>", " ", match.group(1))
        title = html_lib.unescape(re.sub(r"\s+", " ", title)).strip()
        title = re.sub(r"\s*:\s*Amazon\.com\b.*$", "", title, flags=re.I).strip()
        if title:
            return title
    return ""


def rewrite_competitor_title_tail(title, min_chars=100, max_chars=125):
    """Keep the competitor title body and alter only 3–5 trailing words."""
    title = re.sub(r"\s+", " ", str(title or "")).strip(" ,")
    if not title:
        return ""

    additions = (
        "for Everyday Use",
        "for Everyday Project Use",
        "for Organized Everyday Project Use",
    )
    for phrase in additions:
        candidate = f"{title}, {phrase}"
        if min_chars <= len(candidate) <= max_chars:
            return candidate

    words = title.split()
    tail_start = max(0, len(words) - 8)
    deletion_candidates = []
    for count in range(3, 6):
        if len(words) <= count:
            break
        for removed in combinations(range(tail_start, len(words)), count):
            removed = set(removed)
            candidate = " ".join(word for index, word in enumerate(words) if index not in removed).rstrip(" ,-/")
            if min_chars <= len(candidate) <= max_chars:
                deletion_candidates.append((count, -len(candidate), candidate))
    if deletion_candidates:
        return min(deletion_candidates)[2]

    if len(title) < min_chars:
        candidate = f"{title}, {additions[-1]}"
        if len(candidate) <= max_chars:
            return candidate
    return " ".join(words[:-5]).rstrip(" ,-/") if len(words) > 5 else title


def _normalize_competitor_color(value):
    color = re.sub(r"\s+", " ", str(value or "")).strip()
    known_typos = {
        "llight blue": "Light Blue",
    }
    return known_typos.get(color.casefold(), color)


def _resolve_task_price(task, competitor_paths):
    if task.get("price") not in (None, ""):
        return _positive_number(task.get("price"), "price")
    competitor_price = extract_competitor_price(competitor_paths)
    if competitor_price is not None:
        return competitor_price
    if task.get("online_price") not in (None, ""):
        return _positive_number(task.get("online_price"), "online_price")
    if task.get("estimated_price") not in (None, ""):
        return _positive_number(task.get("estimated_price"), "estimated_price")
    raise ValueError(
        "竞品 HTML 没有识别到当前售价；请先检查在线页面并填写 online_price，"
        "在线页面也无价时再填写 estimated_price。"
    )


def _apply_fast_overrides(rows, task, name, price, tier, template_path, copy_defaults=None):
    copy_defaults = copy_defaults or {}
    child_rows = [deepcopy(row) for row in rows if str(row.get("parentage_level") or "").lower() != "parent"]
    if not child_rows:
        raise ValueError(f"{name} 没有提炼出可售子体。")

    variants = task.get("variants") or [{}]
    base_child = child_rows[0]
    normalized_children = []
    for variant_index, variant in enumerate(variants, 1):
        row = deepcopy(base_child)
        row.update({key: value for key, value in variant.items() if value not in (None, "")})
        row["sku"] = _child_sku(task, variant, variant_index)
        row["product_name"] = task.get("product_name") or row.get("product_name") or name
        row["project_name"] = name
        row["route"] = FAST_ROUTE
        row["brand"] = "Generic"
        row["manufacturer"] = "Generic"
        row["parent_sku"] = _parent_sku(task)
        row["parentage_level"] = "Child"
        row["variation_theme"] = _variation_theme(task, template_path)
        row["color"] = variant.get("color") or task.get("color") or row.get("color")
        row["size"] = _normalize_size(variant.get("size") or task.get("size") or row.get("size"))
        row["material"] = variant.get("material") or task.get("material") or row.get("material")
        row["set_count"] = variant.get("set_count") or task.get("set_count") or row.get("set_count") or 1
        row["item_type_keyword"] = task.get("item_type_keyword") or copy_defaults.get("item_type_keyword") or row.get("item_type_keyword")
        row["title"] = (
            variant.get("base_title")
            or task.get("base_title")
            or task.get("_competitor_base_title")
            or row.get("title")
        )
        for field in ["bullet_1", "bullet_2", "bullet_3", "bullet_4", "bullet_5", "description"]:
            row[field] = variant.get(field) or task.get(field) or copy_defaults.get(field) or row.get(field)
        row["list_price"] = variant.get("price") or price
        row["haul_price"] = variant.get("price") or price
        row["package_length_in"] = tier["package_in"][0]
        row["package_width_in"] = tier["package_in"][1]
        row["package_height_in"] = tier["package_in"][2]
        row["package_weight_lb"] = tier["weight_lb"]
        row["country_of_origin"] = "China"
        row["batteries_required"] = "No"
        row["dangerous_goods"] = "Not Applicable"
        row["main_image_url"] = ""
        row["notes"] = f"批量快速上品；物流档位 {tier['label']}；WPS 重算后再上传。"
        row.update(_fast_copy_fallback(row))
        normalized_children.append(row)

    parent = deepcopy(normalized_children[0])
    parent["sku"] = _parent_sku(task)
    parent["parent_sku"] = ""
    parent["parentage_level"] = "Parent"
    parent["manufacturer"] = ""
    parent["color"] = ""
    parent["size"] = ""
    parent["set_count"] = ""
    parent["list_price"] = ""
    parent["haul_price"] = ""
    parent["package_length_in"] = ""
    parent["package_width_in"] = ""
    parent["package_height_in"] = ""
    parent["package_weight_lb"] = ""
    return [parent] + normalized_children


def resolve_logistics_tier(tier_id=None, weight_grams=None, config_path=None):
    config_path = Path(config_path) if config_path else LOGISTICS_TIERS_PATH
    config = json.loads(config_path.read_text(encoding="utf-8"))
    tiers = config["tiers"]
    normalized_id = str(tier_id or "").strip().lower()
    if normalized_id:
        for tier in tiers:
            if normalized_id in {str(tier["id"]).lower(), str(tier["label"]).lower()}:
                return _converted_tier(tier)
        raise ValueError(f"未知物流档位：{tier_id}")

    grams = _positive_number(weight_grams, "weight_grams")
    for tier in tiers:
        if grams <= float(tier["max_weight_grams"]):
            return _converted_tier(tier)
    raise ValueError(f"重量 {grams:g} g 超出当前批量快速路线最高档位。")


def _converted_tier(tier):
    converted = dict(tier)
    converted["package_in"] = [round(float(value) / 2.54, 2) for value in tier["package_cm"]]
    converted["weight_lb"] = round(float(tier["default_weight_grams"]) / 453.592, 2)
    return converted


def _variation_theme(task, template_path):
    requested = str(task.get("variation_theme") or "").strip()
    allowed = _template_variation_themes(template_path)
    if requested:
        if allowed and requested not in allowed:
            raise ValueError(f"Variation Theme {requested} 不在模板 Valid Values 中：{sorted(allowed)}")
        return requested
    for candidate in ("COLOR", "Color", "SIZE", "Size", "SET_NAME", "SetName"):
        if candidate in allowed:
            return candidate
    if allowed:
        return sorted(allowed)[0]
    raise ValueError("模板中未识别到 Variation Theme 有效值，请在批量清单显式填写。")


def _template_variation_themes(path):
    wb = load_workbook(path, data_only=True, read_only=True, keep_vba=Path(path).suffix.lower() == ".xlsm")
    if "Valid Values" not in wb.sheetnames:
        return set()
    ws = wb["Valid Values"]
    product_type = ""
    for row in ws.iter_rows(values_only=True):
        values = [str(value).strip() if value is not None else "" for value in row]
        for position, value in enumerate(values):
            if value.startswith("Product Type"):
                product_type = next((item for item in values[position + 1:] if item and not item.startswith("[")), "")
        if product_type and any(value.startswith(f"Variation Theme Name - [ {product_type} ]") for value in values):
            return {value for value in values if value and not value.startswith("Variation Theme Name")}
    return set()


def _reference_copy(value, base_dir):
    if not value:
        return {}
    path = _resolve_path(value, base_dir)
    wb = load_workbook(path, data_only=True, read_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    ws = wb["Template"] if "Template" in wb.sheetnames else wb.active
    field_to_col = {
        str(ws.cell(5, col).value).strip(): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(5, col).value
    }
    parentage_col = field_to_col.get("parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value")
    sku_col = field_to_col.get("contribution_sku#1.value")
    source_row = None
    for row in range(7, ws.max_row + 1):
        if not sku_col or ws.cell(row, sku_col).value in (None, ""):
            continue
        parentage = str(ws.cell(row, parentage_col).value or "").strip().lower() if parentage_col else ""
        if parentage != "parent":
            source_row = row
            break
    if source_row is None:
        return {}
    fields = {
        "item_type_keyword": "item_type_keyword[marketplace_id=ATVPDKIKX0DER]#1.value",
        "bullet_1": "bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
        "bullet_2": "bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value",
        "bullet_3": "bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#3.value",
        "bullet_4": "bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#4.value",
        "bullet_5": "bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#5.value",
        "description": "product_description[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    }
    return {
        key: ws.cell(source_row, field_to_col[field_name]).value
        for key, field_name in fields.items()
        if field_name in field_to_col and ws.cell(source_row, field_to_col[field_name]).value not in (None, "")
    }


def _reference_safe_fields(value, base_dir):
    if not value:
        return {}
    path = _resolve_path(value, base_dir)
    wb = load_workbook(path, data_only=True, read_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    ws = wb["Template"] if "Template" in wb.sheetnames else wb.active
    field_to_col = {
        str(ws.cell(5, col).value).strip(): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(5, col).value
    }
    parentage_col = field_to_col.get("parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value")
    sku_col = field_to_col.get("contribution_sku#1.value")
    source_row = None
    for row in range(7, ws.max_row + 1):
        if not sku_col or ws.cell(row, sku_col).value in (None, ""):
            continue
        parentage = str(ws.cell(row, parentage_col).value or "").strip().lower() if parentage_col else ""
        if parentage != "parent":
            source_row = row
            break
    if source_row is None:
        return {}
    return {
        field_name: ws.cell(source_row, field_to_col[field_name]).value
        for field_name in REFERENCE_SAFE_FIELDS
        if field_name in field_to_col and ws.cell(source_row, field_to_col[field_name]).value not in (None, "")
    }


def _reference_rows(value, base_dir):
    if not value:
        return {}
    path = _resolve_path(value, base_dir)
    wb = load_workbook(path, data_only=True, read_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    ws = wb["Template"] if "Template" in wb.sheetnames else wb.active
    field_to_col = {
        str(ws.cell(5, col).value).strip(): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(5, col).value
    }
    sku_col = field_to_col.get("contribution_sku#1.value")
    if not sku_col:
        return {}
    rows = {}
    for row in range(7, ws.max_row + 1):
        sku = str(ws.cell(row, sku_col).value or "").strip()
        if not sku:
            continue
        rows[sku] = {
            field_name: ws.cell(row, col).value
            for field_name, col in field_to_col.items()
            if field_name != "::record_action"
        }
    return rows


def _apply_batch_overlay(path, task, tier, reference_fields, reference_rows=None):
    path = Path(path)
    wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    ws = wb["Template"] if "Template" in wb.sheetnames else wb.active
    field_to_col = {
        str(ws.cell(5, col).value).strip(): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(5, col).value
    }
    sku_col = field_to_col.get("contribution_sku#1.value")
    parentage_col = field_to_col.get("parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value")
    extra_fields = task.get("extra_fields") or {}
    dimension_fields = _dimension_overlay_fields(tier)
    protected_blank_tokens = (
        "image",
        "media_location",
        "minimum_seller_allowed_price",
        "maximum_seller_allowed_price",
        "skip_offer",
    )
    reference_rows = reference_rows or {}
    requested_action = str(task.get("listing_action") or "").strip()
    if requested_action.casefold() == "edit":
        requested_action = "Edit (Partial Update)"

    for row in range(7, ws.max_row + 1):
        if not sku_col or ws.cell(row, sku_col).value in (None, ""):
            continue
        parentage = str(ws.cell(row, parentage_col).value or "").strip().lower() if parentage_col else ""
        is_parent = parentage == "parent"
        for field_name, col in field_to_col.items():
            lowered = field_name.lower()
            if any(token in lowered for token in protected_blank_tokens):
                ws.cell(row, col).value = None
        if is_parent:
            action_col = field_to_col.get("::record_action")
            if action_col and requested_action:
                ws.cell(row, action_col).value = requested_action
            elif action_col:
                sku = str(ws.cell(row, sku_col).value or "").strip()
                previous = reference_rows.get(sku)
                if previous and _row_changed(ws, row, field_to_col, previous):
                    ws.cell(row, action_col).value = "Edit (Partial Update)"
            continue
        for field_name, value in {**reference_fields, **extra_fields, **dimension_fields}.items():
            col = field_to_col.get(field_name)
            if not col or value in (None, ""):
                continue
            if any(token in field_name.lower() for token in ("image", "media_location")):
                continue
            if field_name in reference_fields and ws.cell(row, col).value not in (None, ""):
                continue
            ws.cell(row, col).value = value
        action_col = field_to_col.get("::record_action")
        if action_col and requested_action:
            ws.cell(row, action_col).value = requested_action
        elif action_col:
            sku = str(ws.cell(row, sku_col).value or "").strip()
            previous = reference_rows.get(sku)
            if previous and _row_changed(ws, row, field_to_col, previous):
                ws.cell(row, action_col).value = "Edit (Partial Update)"
    wb.save(path)


def _row_changed(ws, row, field_to_col, previous):
    ignored = {"::record_action"}
    for field_name, col in field_to_col.items():
        if field_name in ignored or field_name not in previous:
            continue
        current_value = ws.cell(row, col).value
        previous_value = previous[field_name]
        if _comparable_value(current_value) != _comparable_value(previous_value):
            return True
    return False


def _comparable_value(value):
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _dimension_overlay_fields(tier):
    length, width, height = tier["package_in"]
    weight = tier["weight_lb"]
    fields = {}
    for prefix in ["item_dimensions", "item_package_dimensions", "item_length_width_height"]:
        fields[f"{prefix}[marketplace_id=ATVPDKIKX0DER]#1.length.value"] = length
        fields[f"{prefix}[marketplace_id=ATVPDKIKX0DER]#1.length.unit"] = "Inches"
        fields[f"{prefix}[marketplace_id=ATVPDKIKX0DER]#1.width.value"] = width
        fields[f"{prefix}[marketplace_id=ATVPDKIKX0DER]#1.width.unit"] = "Inches"
        fields[f"{prefix}[marketplace_id=ATVPDKIKX0DER]#1.height.value"] = height
        fields[f"{prefix}[marketplace_id=ATVPDKIKX0DER]#1.height.unit"] = "Inches"
    fields.update({
        "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.height.value": length,
        "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.height.unit": "Inches",
        "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.width.value": width,
        "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.width.unit": "Inches",
        "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.depth.value": height,
        "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.depth.unit": "Inches",
    })
    for prefix in ["item_weight", "item_package_weight", "item_display_weight"]:
        fields[f"{prefix}[marketplace_id=ATVPDKIKX0DER]#1.value"] = weight
        fields[f"{prefix}[marketplace_id=ATVPDKIKX0DER]#1.unit"] = "Pounds"
    return fields


def _parent_sku(task):
    value = str(task.get("parent_sku") or "").strip()
    if value:
        return value
    base = _sku_token(task.get("sku_base") or task.get("name") or task.get("output_name"))
    return f"CA-{base}"


def _child_sku(task, variant, index):
    explicit = str(variant.get("sku") or (task.get("child_sku") if index == 1 else "") or "").strip()
    if explicit:
        return explicit
    parts = [_parent_sku(task)]
    color = _sku_token(variant.get("color") or task.get("color"))
    size = _sku_token(variant.get("size") or task.get("size"))
    suffix = f"{color}{size}" or f"V{index}"
    parts.append(suffix)
    return "-".join(parts)


def _sku_token(value):
    return "".join(re.findall(r"[A-Z0-9]+", str(value or "").upper()))[:30] or "PRODUCT"


def _normalize_size(value):
    text = str(value or "").strip()
    return re.sub(r"\bmm\b", "MM", text, flags=re.I)


def _fast_copy_fallback(row):
    subject = str(row.get("product_name") or "General Product").strip()
    color = str(row.get("color") or "Selected Color").strip()
    size = str(row.get("size") or "Selected Size").strip()
    material = str(row.get("material") or "the listed material").strip()
    count = str(row.get("set_count") or 1).strip()

    title = str(row.get("title") or "").strip()
    if len(title) < 100:
        title = (
            f"{subject}, Practical Supplies for Home, Travel, Work, Craft Projects, "
            "Daily Tasks and General Organization"
        )

    bullet_sources = [
        f"Package Details Included: This listing contains {count} count of {subject} in {color} and {size}, giving buyers a clearly identified configuration for planned projects, routine tasks, organized storage, and general everyday use.",
        f"Material And Finish Notes: The item is listed with {material} construction and a {color} appearance, while the selected {size} format helps buyers compare the configuration with their intended use before ordering.",
        f"Flexible Project Planning: Use the selected {subject} configuration for suitable home, work, travel, craft, organization, or general project needs, and combine it with separately purchased tools or accessories when required.",
        f"Simple Handling And Storage: The compact product format supports straightforward placement, sorting, carrying, and storage between uses, while the clearly listed color, size, and count help keep supplies organized.",
        f"Review Before Ordering: Confirm the selected {color} color, {size} size, {count} count, and listed package details before purchase so the chosen configuration matches the intended task, project plan, or storage arrangement.",
    ]
    bullets = {}
    for index, source in enumerate(bullet_sources, 1):
        current = str(row.get(f"bullet_{index}") or "").strip()
        bullets[f"bullet_{index}"] = current if 200 <= len(current) <= 250 else _fit_bullet(source)

    description = str(row.get("description") or "").strip()
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n+", description) if part.strip()]
    if not (1500 <= len(description) <= 1800 and len(paragraphs) == 4):
        description = _fast_description(subject, color, size, material, count)

    return {"title": title, "description": description, **bullets}


def _fit_bullet(text):
    additions = [
        " Review the listing details before ordering to confirm the selected configuration.",
        " Keep the pieces together between uses for simple sorting and project preparation.",
        " Additional tools, containers, and accessories are not included unless stated.",
    ]
    value = text.strip()
    addition_index = 0
    while len(value) < 200:
        value += additions[addition_index % len(additions)]
        addition_index += 1
    if len(value) > 250:
        value = value[:249].rsplit(" ", 1)[0]
    return value.rstrip(" .") + "."


def _fast_description(subject, color, size, material, count):
    paragraphs = [
        f"Product Overview: This listing presents {subject} in the selected {color} color and {size} size. The package contains {count} count as stated in the product configuration. The item is intended for suitable everyday projects, organization, work areas, travel preparation, craft activities, or other ordinary uses that match the product type. Review the selected variation before ordering so the received configuration aligns with the planned task.",
        f"Material And Configuration: The product is listed with {material} construction. Color, size, count, and package information are provided as variation details to help distinguish this option from other configurations in the same family. Surface appearance and small dimensional differences may be viewed under different lighting or handling conditions. Compare the listed attributes with the intended application before use.",
        f"Use And Organization: Arrange the included pieces according to the planned project, workspace, storage method, or general routine. The product can be kept with related supplies so the selected configuration is easier to identify between uses. Tools, containers, finished projects, display items, and unrelated accessories are not included unless they are specifically named in the package contents. Follow ordinary handling appropriate for the product type.",
        f"Package Review: The shipment is based on the selected {color}, {size}, and {count} count option shown in this listing. Check the package contents after receipt and keep the product information for future reference. Store the item in a suitable location between uses and keep separate components organized when applicable. This description focuses on the supplied product configuration and does not include promotional, service, certification, or performance promises.",
    ]
    padding = " Review the listed variation, package count, and product details before use so the configuration remains easy to identify and organize."
    while len("\n\n".join(paragraphs)) < 1500:
        shortest = min(range(4), key=lambda index: len(paragraphs[index]))
        paragraphs[shortest] += padding
    return "\n\n".join(paragraphs)


def _output_filename(task, name, template_suffix):
    requested = str(task.get("output_name") or name).strip()
    requested_path = Path(requested)
    requested_suffix = requested_path.suffix.lower()
    suffix = requested_suffix if requested_suffix in {".xlsx", ".xlsm"} else template_suffix
    suffix = suffix if suffix.lower() in {".xlsx", ".xlsm"} else ".xlsx"
    requested_stem = requested_path.stem if requested_suffix in {".xlsx", ".xlsm"} else requested
    stem = requested_stem if re.search(r"V\d+$", requested_stem, re.I) else f"{requested_stem}V1"
    return f"{safe_name(stem)}{suffix}"


def _competitor_paths(task, base_dir):
    values = task.get("competitors") or task.get("competitor") or []
    if isinstance(values, str):
        values = [values]
    paths = [_resolve_path(value, base_dir) for value in values]
    if not paths:
        raise ValueError("缺少 competitor/competitors 竞品HTML路径。")
    return paths


def _required_path(task, field, base_dir):
    value = task.get(field)
    if not value:
        raise ValueError(f"缺少 {field}。")
    return _resolve_path(value, base_dir)


def _resolve_optional_path(value, base_dir):
    return _resolve_path(value, base_dir, must_exist=False) if value else None


def _resolve_path(value, base_dir, must_exist=True):
    path = Path(str(value)).expanduser()
    if not path.is_absolute():
        path = Path(base_dir) / path
    path = path.resolve()
    if must_exist and not path.exists():
        raise FileNotFoundError(path)
    return path


def _positive_number(value, label):
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label} 必须是正数。") from None
    if number <= 0:
        raise ValueError(f"{label} 必须是正数。")
    return number


def _check_zip_structure(path):
    if Path(path).suffix.lower() not in {".xlsx", ".xlsm"}:
        return
    if not zipfile.is_zipfile(path):
        raise ValueError(f"输出文件不是有效的 Excel 压缩结构：{path}")
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise ValueError(f"Excel 压缩结构损坏：{bad_member}")
