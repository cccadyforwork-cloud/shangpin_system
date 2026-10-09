"""Persistent data model for the eight-step batch listing workbench.

This module is intentionally an orchestration layer.  It does not replace or
modify the Amazon template writer/validator.  Generated Amazon workbooks must
still pass through the existing deterministic core before delivery.
"""

import io
import json
import re
import shutil
import threading
from datetime import datetime
from functools import wraps
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .batch_fast_prelisting import _template_variation_themes
from .paths import DATA_DIR, ROOT, safe_name
from .template_sheet import find_template_sheet
from .template_writer import extract_template_product_type


WORKBENCH_DIR = DATA_DIR / "automation_workbench"
LEDGER_PATH = WORKBENCH_DIR / "ledger.json"
FILES_DIR = WORKBENCH_DIR / "files"
EXPORTS_DIR = WORKBENCH_DIR / "exports"
MANUAL_WORKBENCH_DIR = DATA_DIR / "manual_listing_workbench"
MANUAL_LEDGER_PATH = MANUAL_WORKBENCH_DIR / "ledger.json"
MANUAL_FILES_DIR = MANUAL_WORKBENCH_DIR / "files"
MANUAL_EXPORTS_DIR = MANUAL_WORKBENCH_DIR / "exports"

EDITABLE_FIELDS = {
    "competitor_url",
    "variation_theme",
    "procurement_url",
    "procurement_name",
    "procurement_quantity",
    "purchase_cost",
    "list_price",
    "haul_price",
    "supplier_name",
    "tax_unit_price",
    "msku",
    "product_name",
    "optimized_title",
    "bullet_1",
    "bullet_2",
    "bullet_3",
    "bullet_4",
    "bullet_5",
    "note",
}
_LOCK = threading.RLock()


def _synchronized(function):
    @wraps(function)
    def wrapped(*args, **kwargs):
        with _LOCK:
            return function(*args, **kwargs)

    return wrapped


def _now():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def canonical_amazon_url(value):
    text = str(value or "").strip()
    match = re.search(r"(?:/dp/|[?&]ASIN=)([A-Z0-9]{10})", text, re.I)
    if not match:
        return text, ""
    asin = match.group(1).upper()
    return f"https://www.amazon.com/dp/{asin}", asin


def validate_haul_template(path):
    """Reject ordinary Amazon catalog workbooks before they enter the fill flow."""
    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=False, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        sheet = find_template_sheet(workbook)
        if sheet is None:
            raise ValueError("这不是可识别的 Amazon 模板。")
        haul_marker = False
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True):
            if any("[audience=BZR]" in str(value or "") for value in row):
                haul_marker = True
                break
        if not haul_marker:
            raise ValueError("已拦截主站模板：请从 Seller Central“下载亚马逊 Haul 电子表格”入口重新下载。")
    finally:
        workbook.close()
    return True


def _empty_ledger():
    return {"version": 1, "updated_at": "", "batches": []}


def load_ledger(path=LEDGER_PATH):
    path = Path(path)
    if not path.exists():
        return _empty_ledger()
    data = json.loads(path.read_text(encoding="utf-8"))
    data.setdefault("version", 1)
    data.setdefault("updated_at", "")
    data.setdefault("batches", [])
    for batch in data["batches"]:
        for row in batch.get("rows", []):
            row.setdefault("variant_rows", [])
            analysis = row.setdefault("analysis", {})
            if analysis.get("phase") == "variation_ready":
                count = analysis.get("variant_count") or len(row["variant_rows"])
                analysis["phase"] = "fill_queued"
                analysis["progress"] = 45
                analysis["eta_label"] = "执行器接单后预计 10–20 分钟"
                analysis["phase_label"] = (
                    f"进度 45%：已保存 {count} 个变体；自动填表排队中，尚未开始写表；"
                    "执行器接单后预计 10–20 分钟"
                )
    return data


def variation_dimensions(theme):
    """Return exact dimension names represented by a template variation theme."""
    value = str(theme or "").split(" (Deprecated:", 1)[0].strip()
    return [part.strip() for part in value.split("/") if part.strip()]


def _clean_variant_rows(theme, rows):
    dimensions = variation_dimensions(theme)
    if not dimensions or not isinstance(rows, list):
        return []
    cleaned = []
    for raw in rows[:100]:
        if not isinstance(raw, dict):
            continue
        item = {dimension: str(raw.get(dimension) or "").strip()[:200] for dimension in dimensions}
        if any(item.values()):
            if "NUMBER_OF_ITEMS" in item and item["NUMBER_OF_ITEMS"]:
                if not item["NUMBER_OF_ITEMS"].isdigit() or int(item["NUMBER_OF_ITEMS"]) <= 0:
                    raise ValueError("数量必须是大于 0 的整数。")
            cleaned.append(item)
    return cleaned


@_synchronized
def save_ledger(data, path=LEDGER_PATH):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    data["updated_at"] = _now()
    temporary = path.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)
    return path


def ensure_ledger(path=LEDGER_PATH):
    return load_ledger(path)


@_synchronized
def add_manual_product(competitor_url, owner="自主", ledger_path=MANUAL_LEDGER_PATH):
    link, asin = canonical_amazon_url(competitor_url)
    if not link:
        raise ValueError("请填写竞品链接。")
    data = load_ledger(ledger_path)
    if data.get("batches"):
        batch = data["batches"][0]
    else:
        batch = {"id": "manual-listing", "name": "自主上新品", "source_name": "逐条新增", "owners": [], "rows": []}
        data["batches"] = [batch]
    owner = str(owner or "自主").strip() or "自主"
    identity = asin or f"ROW-{len(batch['rows']) + 1}"
    existing = next((row for row in batch["rows"] if row.get("owner") == owner and (row.get("asin") or row.get("id")) == identity), None)
    if existing:
        return existing
    row = _new_row(identity, asin, link, owner, len(batch["rows"]) + 2)
    _apply_experience_hits([row])
    batch["rows"].insert(0, row)
    batch["owners"] = list(dict.fromkeys([owner] + list(batch.get("owners") or [])))
    save_ledger(data, ledger_path)
    return row


@_synchronized
def import_selection_sheet(source_path, ledger_path=LEDGER_PATH):
    source = Path(source_path).expanduser().resolve()
    if not source.is_file() or source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("请选择有效的选品 Excel 文件。")
    data = load_ledger(ledger_path)
    previous = {}
    for batch in data.get("batches", []):
        for row in batch.get("rows", []):
            previous[(str(row.get("owner") or ""), str(row.get("asin") or row.get("id") or ""))] = row

    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=source.suffix.lower() == ".xlsm")
    sheet = workbook.active
    headers = {str(sheet.cell(1, col).value or "").strip(): col for col in range(1, sheet.max_column + 1)}
    link_col = headers.get("LINK") or headers.get("链接") or headers.get("竞品链接") or 1
    owner_col = headers.get("负责人") or 2
    rows = []
    seen = set()
    for row_number in range(2, sheet.max_row + 1):
        raw_link = str(sheet.cell(row_number, link_col).value or "").strip()
        if not raw_link:
            continue
        link, asin = canonical_amazon_url(raw_link)
        owner = str(sheet.cell(row_number, owner_col).value or "待分配").strip() or "待分配"
        identity = asin or f"ROW-{row_number}"
        key = (owner, identity)
        if key in seen:
            continue
        seen.add(key)
        old = previous.get(key, {})
        row = _new_row(identity, asin, link, owner, row_number)
        for field in EDITABLE_FIELDS:
            if field in old:
                row[field] = old[field]
        for field in ("analysis", "source_template", "template_info", "image_job", "lingxing_file", "amazon_output", "sync_status"):
            if field in old:
                row[field] = old[field]
        rows.append(row)
    workbook.close()

    batch_id = safe_name(source.stem).lower()
    batch = {
        "id": batch_id,
        "name": source.stem,
        "source_name": source.name,
        "owners": list(dict.fromkeys(row["owner"] for row in rows)),
        "rows": rows,
    }
    _apply_experience_hits(rows)
    other_batches = [item for item in data.get("batches", []) if item.get("id") != batch_id]
    data["batches"] = [batch] + other_batches
    save_ledger(data, ledger_path)
    return batch


@_synchronized
def refresh_experience_hits(ledger_path=LEDGER_PATH):
    data = load_ledger(ledger_path)
    rows = [row for batch in data.get("batches", []) for row in batch.get("rows", [])]
    _apply_experience_hits(rows)
    save_ledger(data, ledger_path)
    return sum(1 for row in rows if row.get("analysis", {}).get("experience_hits"))


def _apply_experience_hits(rows):
    try:
        from .experience_pool import lookup_many_asins

        matches = lookup_many_asins([row.get("asin") for row in rows])
    except Exception:
        matches = {}
    for row in rows:
        row.setdefault("analysis", {})["experience_hits"] = matches.get(str(row.get("asin") or "").upper(), [])


def _new_row(identity, asin, link, owner, row_number):
    return {
        "id": f"{identity}--{safe_name(owner)}",
        "asin": asin,
        "row_number": row_number,
        "owner": owner,
        "competitor_url": link,
        "analysis": {
            "status": "待分析",
            "original_title": "",
            "traffic_coverage": None,
            "keywords": [],
            "image_reference": "",
            "experience_hits": [],
        },
        "source_template": "",
        "template_info": {
            "product_type": "",
            "variation_themes": [],
            "status": "待从 Seller Central Haul 入口下载",
            "download_route": "批量上传商品 > 下载空白模板 > 在亚马逊 Haul 中发布商品 > 下载电子表格",
        },
        "variation_theme": "",
        "variant_rows": [],
        "procurement_url": "",
        "procurement_name": "",
        "procurement_quantity": "10",
        "purchase_cost": "",
        "list_price": "",
        "haul_price": "",
        "supplier_name": "",
        "tax_unit_price": "",
        "msku": "",
        "product_name": "",
        "optimized_title": "",
        "bullet_1": "",
        "bullet_2": "",
        "bullet_3": "",
        "bullet_4": "",
        "bullet_5": "",
        "image_job": {"status": "未开始", "progress": 0, "folder_name": "", "prompts": []},
        "amazon_output": "",
        "lingxing_file": "",
        "sync_status": "等待生成 Amazon 模板",
        "note": "",
    }


@_synchronized
def update_row(batch_id, row_id, fields, ledger_path=LEDGER_PATH):
    data = load_ledger(ledger_path)
    row = _find_row(data, batch_id, row_id)
    changed_identity = False
    old_theme = row.get("variation_theme", "")
    for key, value in (fields or {}).items():
        if key == "variant_rows":
            row["variant_rows"] = _clean_variant_rows(row.get("variation_theme"), value)
            analysis = row.setdefault("analysis", {})
            dimensions = variation_dimensions(row.get("variation_theme"))
            complete = sum(all(item.get(dimension) for dimension in dimensions) for item in row["variant_rows"])
            analysis["variant_count"] = complete
            if complete:
                analysis["status"] = "分析中"
                analysis["phase"] = "fill_queued"
                analysis["phase_label"] = (
                    f"进度 45%：已保存 {complete} 个变体；自动填表排队中，尚未开始写表；"
                    "执行器接单后预计 10–20 分钟"
                )
                analysis["progress"] = 45
                analysis["eta_label"] = "执行器接单后预计 10–20 分钟"
                analysis["queued_at"] = _now()
            continue
        if key not in EDITABLE_FIELDS:
            continue
        text = str(value or "").strip()
        if key == "competitor_url":
            text, asin = canonical_amazon_url(text)
            if asin:
                row["asin"] = asin
        if key == "procurement_quantity" and text:
            try:
                if float(text) <= 0:
                    raise ValueError
            except ValueError as exc:
                raise ValueError("采购数量必须大于 0。") from exc
        row[key] = text
        changed_identity = changed_identity or key in {"msku", "product_name", "variation_theme"}
    if row.get("variation_theme") != old_theme:
        row["variant_rows"] = []
    if changed_identity:
        row["sync_status"] = "信息已修改，生成模板时自动同步"
        row["lingxing_file"] = ""
        row["image_job"]["folder_name"] = _image_folder_name(row)
    save_ledger(data, ledger_path)
    return row


@_synchronized
def claim_analysis(batch_id, row_id, ledger_path=LEDGER_PATH):
    """Queue one visible row for the linked analysis worker."""
    data = load_ledger(ledger_path)
    row = _find_row(data, batch_id, row_id)
    had_source_template = bool(row.get("source_template"))
    analysis = row.setdefault("analysis", {})
    if analysis.get("status") == "分析完成":
        raise ValueError("该产品已完成分析。")
    analysis["status"] = "分析中"
    analysis["phase"] = "experience_lookup"
    analysis["phase_label"] = "正在检索上品经验池"
    analysis["requested_at"] = _now()
    row.setdefault("template_info", {})["status"] = "等待经验池确认 Product Type"
    save_ledger(data, ledger_path)
    return row


@_synchronized
def claim_all_analysis(batch_id, ledger_path=LEDGER_PATH):
    """Claim every unfinished analysis row in one manual or batch route."""
    data = load_ledger(ledger_path)
    batch = _find_batch(data, batch_id)
    count = 0
    claimed_at = _now()
    for row in batch.get("rows", []):
        analysis = row.setdefault("analysis", {})
        if analysis.get("status") in {"分析中", "分析完成"}:
            continue
        analysis["status"] = "分析中"
        analysis["phase"] = "experience_lookup"
        analysis["phase_label"] = "正在检索上品经验池"
        analysis["requested_at"] = claimed_at
        row.setdefault("template_info", {})["status"] = "等待经验池确认 Product Type"
        count += 1
    save_ledger(data, ledger_path)
    return count


@_synchronized
def attach_template(batch_id, row_id, filename, source, ledger_path=LEDGER_PATH, files_dir=FILES_DIR):
    suffix = Path(filename or "").suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Amazon 模板只支持 xlsx 或 xlsm。")
    data = load_ledger(ledger_path)
    row = _find_row(data, batch_id, row_id)
    target_dir = Path(files_dir) / safe_name(batch_id) / safe_name(row_id) / "source_template"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / safe_name(Path(filename).name)
    with target.open("wb") as output:
        shutil.copyfileobj(source, output)
    try:
        validate_haul_template(target)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    product_type = extract_template_product_type(target)
    themes = sorted(_template_variation_themes(target))
    row["source_template"] = _store_path(target)
    row["template_info"] = {
        "product_type": product_type,
        "variation_themes": themes,
        "status": "模板已识别",
    }
    # Users may choose a theme before the Haul workbook is downloaded. Keep
    # that early choice when this exact Product Type supports it; otherwise
    # clear it so an invalid cross-category value cannot reach the writer.
    if row.get("variation_theme") and row.get("variation_theme") not in themes:
        row["variation_theme"] = ""
        row["variant_rows"] = []
    elif not row.get("variation_theme"):
        row["variation_theme"] = themes[0] if len(themes) == 1 else ""
    row["sync_status"] = "模板已就绪，等待分析和生成"
    save_ledger(data, ledger_path)
    return target, row


@_synchronized
def prepare_image_job(batch_id, row_id, ledger_path=LEDGER_PATH):
    data = load_ledger(ledger_path)
    row = _find_row(data, batch_id, row_id)
    if not row.get("product_name") and not row.get("optimized_title"):
        raise ValueError("请先填写产品名或完成标题分析。")
    prompts = build_image_prompts(row)
    row["image_job"] = {
        "status": "待上品电脑领取",
        "progress": 0,
        "folder_name": _image_folder_name(row),
        "prompts": prompts,
    }
    save_ledger(data, ledger_path)
    return row["image_job"]


def build_image_prompts(row):
    product = str(row.get("product_name") or row.get("optimized_title") or "the product").strip()
    source = str(row.get("competitor_url") or "")
    common = (
        f"Use the supplied reference photos for {product} from {source}. Preserve the exact product shape, material, color, "
        "parts and included quantity. Photorealistic Amazon Haul product photography, consistent visual style, high resolution. "
        "Any visible text must be minimal English only. No Chinese, price, discount, brand watermark, platform logo, medical claim, "
        "exaggerated promise or unverified accessory."
    )
    return [
        common + " Create a natural everyday lifestyle scene showing the product in its real use, with soft realistic light and clear buyer context. No text.",
        common + " Create a clean core-feature image with one clear hero product and two or three inset close-ups. Use only short verified labels such as Easy to Use, Durable Material, or Compact Design when accurate.",
        common + " Create a vertical multi-scene poster. Place four realistic capsule-shaped use-scene panels across the lower area and a product or short heading such as Everyday Use above.",
        common + " Create a premium material-and-detail image over a softly faded real-use background. Show close-ups of verified material, structure, handling or portability with two or three short factual labels. Do not show dimensions or a parameter table.",
        common + " Create an Amazon-compliant main image on pure white (#FFFFFF). Show only the exact sold product, centered, fully visible, filling about 85% of frame. No text, hands, props, scenes, dimensions, icons or decoration.",
        common + " Create either a four-step use/instruction infographic or a factual specification/dimension image, choosing the format that best fits this exact product. Use clear English steps or verified measurements only; never invent dimensions, functions, parts or accessories.",
    ]


def export_lingxing(batch_id, owner="", ledger_path=LEDGER_PATH, exports_dir=EXPORTS_DIR):
    data = load_ledger(ledger_path)
    batch = _find_batch(data, batch_id)
    selected = [row for row in batch.get("rows", []) if (not owner or row.get("owner") == owner)]
    selected = [row for row in selected if row.get("msku") or row.get("product_name")]
    if not selected:
        raise ValueError("请先填写至少一行 MSKU 或产品名。")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品"
    headers = ["*SKU", "品名", "型号", "关联数量1", "采购员", "采购成本(CNY)", "供应商名称", "含税单价"]
    sheet.append(headers)
    for row in selected:
        sheet.append([
            row.get("msku", ""),
            row.get("product_name", ""),
            row.get("msku", ""),
            _number(row.get("procurement_quantity")),
            row.get("owner", ""),
            _number(row.get("purchase_cost")),
            row.get("supplier_name", ""),
            _number(row.get("tax_unit_price")),
        ])
    header_fill = PatternFill("solid", fgColor="4472C4")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    widths = [24, 28, 24, 14, 14, 18, 24, 14]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    target_dir = Path(exports_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    suffix = f"_{safe_name(owner)}" if owner else ""
    target = target_dir / f"{safe_name(batch.get('name') or batch_id)}{suffix}_领星上传.xlsx"
    workbook.save(target)
    workbook.close()
    relative = _store_path(target)
    for row in selected:
        row["lingxing_file"] = relative
    save_ledger(data, ledger_path)
    return target, len(selected)


def payload(ledger_path=LEDGER_PATH):
    data = load_ledger(ledger_path)
    batches = []
    totals = {"rows": 0, "templates": 0, "analyzed": 0, "images": 0}
    for batch in data.get("batches", []):
        rendered_rows = []
        for row in batch.get("rows", []):
            rendered = dict(row)
            rendered["template_file"] = _file_info(row.get("source_template"))
            rendered["lingxing"] = _file_info(row.get("lingxing_file"))
            rendered["amazon_file"] = _file_info(row.get("amazon_output"))
            rendered_rows.append(rendered)
            totals["rows"] += 1
            totals["templates"] += int(rendered["template_file"]["exists"])
            totals["analyzed"] += int(row.get("analysis", {}).get("status") == "分析完成")
            totals["images"] += int(row.get("image_job", {}).get("status") == "生成完成")
        batches.append({**batch, "rows": rendered_rows})
    return {"ok": True, "updated_at": data.get("updated_at", ""), "totals": totals, "batches": batches}


def referenced_file(value, ledger_path=LEDGER_PATH):
    raw = str(value or "").strip()
    if not raw:
        return None
    candidate = _resolve_path(raw)
    allowed = set()
    data = load_ledger(ledger_path)
    for batch in data.get("batches", []):
        for row in batch.get("rows", []):
            for field in ("source_template", "lingxing_file", "amazon_output"):
                if row.get(field):
                    allowed.add(str(_resolve_path(row[field])))
    return candidate if str(candidate) in allowed else None


def _find_batch(data, batch_id):
    batch = next((item for item in data.get("batches", []) if item.get("id") == batch_id), None)
    if batch is None:
        raise ValueError("找不到批次。")
    return batch


def _find_row(data, batch_id, row_id):
    batch = _find_batch(data, batch_id)
    row = next((item for item in batch.get("rows", []) if item.get("id") == row_id), None)
    if row is None:
        raise ValueError("找不到产品行。")
    return row


def _resolve_path(value):
    path = Path(str(value or ""))
    return path.resolve() if path.is_absolute() else (ROOT / path).resolve()


def _store_path(value):
    path = _resolve_path(value)
    try:
        return path.relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return str(path)


def _file_info(value):
    raw = str(value or "").strip()
    path = _resolve_path(raw) if raw else None
    exists = bool(path and path.is_file())
    return {"exists": exists, "path": str(path) if path else "", "name": path.name if path else ""}


def _image_folder_name(row):
    return safe_name(f"{row.get('msku', '')}{row.get('product_name', '')}") if (row.get("msku") or row.get("product_name")) else ""


def _number(value):
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return text
