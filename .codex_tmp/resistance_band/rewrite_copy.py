from pathlib import Path
from shutil import copyfile
import sys

from openpyxl import load_workbook

sys.path.insert(0, "/Users/cc/Documents/GitHub/shangpin_system")
from app.listing_rules import validate_listing_row


SOURCE = Path("/Users/cc/Documents/GitHub/shangpin_system/outputs/01a09a0d-5e1d-7473-8611-f221c4c977bf/弹力带价格刷新V1.xlsx")
OUTPUT = Path("/Users/cc/Documents/GitHub/shangpin_system/outputs/01a09a0d-5e1d-7473-8611-f221c4c977bf/弹力带价格刷新V2.xlsx")

DESCRIPTION = """This flat resistance band is designed for yoga, Pilates, stretching, warmups, mobility work, and general strength exercises. Its long rectangular profile supports a range of controlled upper-body and lower-body movements without taking up the space of larger exercise equipment. The band can be used at home, in a fitness studio, at the gym, or while traveling.

Thermoplastic elastomer gives the band a smooth surface and flexible feel. The material stretches across its length during use and returns to a compact form for storage. Each band measures approximately 59.06 inches long and 5.91 inches wide. Thickness varies by color option, so customers can select the version that matches the option shown in the listing. The wide, flat shape provides a broad contact area for gripping and positioning during controlled movements.

The flat format can be incorporated into warmup sequences, yoga positions, Pilates movements, mobility drills, and resistance routines for the arms, shoulders, back, legs, and hips. Users can adjust hand placement, grip distance, and stretch range to fit the selected movement. The broad band surface supports repeated pulls, static holds, controlled extensions, and gradual movement sequences. The package contains one exercise resistance band in the selected color option.

For routine care, wipe the band with a soft, slightly damp cloth and allow it to air dry completely before storage. Keep it away from sharp edges, rough surfaces, excessive heat, and prolonged direct sunlight. Roll or fold it loosely and place it in a drawer, gym bag, or storage case. Inspect the surface before each session and discontinue use if the band shows cuts, cracks, or other damage."""

COMMON_BULLETS = [
    "Flexible Band Construction: Thermoplastic elastomer provides a smooth surface and consistent stretch for controlled exercise movements, including warmups, mobility drills, yoga positions, Pilates sessions, and general strength routines.",
    "Portable Workout Format: The flat band rolls or folds into a compact shape for storage in a gym bag, drawer, or travel case, making it convenient to carry between home, fitness studio, hotel, and outdoor training locations.",
    "Versatile Training Uses: Use the band for upper-body and lower-body movements, stretching sequences, warmups, yoga, Pilates, and resistance exercises, with a simple flat design that works for a range of everyday fitness routines.",
    None,
    "Simple Care and Storage: Wipe the band with a soft, slightly damp cloth after use, allow it to air dry fully, and store it away from sharp edges, excessive heat, and direct sunlight to help maintain its surface and shape between sessions.",
]

DIMENSION_BULLETS = {
    "CA-LLD-Single": "Practical Band Dimensions: Each color option measures approximately 59.06 inches long and 5.91 inches wide, while thickness varies by option, providing a broad contact area for controlled holds, extensions, and repeated movements.",
    "LLD-TPE-BLUE": "Practical Band Dimensions: The blue option measures approximately 59.06 inches long and 5.91 inches wide, with a thickness of about 0.014 inch, providing a broad contact area for controlled holds, extensions, and repeated movements.",
    "LLD-TPE-GREEN": "Practical Band Dimensions: The green option measures approximately 59.06 inches long and 5.91 inches wide, with a thickness of about 0.014 inch, providing a broad contact area for controlled holds, extensions, and repeated movements.",
    "LLD-TPE-PINK": "Practical Band Dimensions: The pink option measures approximately 59.06 inches long and 5.91 inches wide, with a thickness of about 0.015 inch, providing a broad contact area for controlled holds, extensions, and repeated movements.",
}


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
copyfile(SOURCE, OUTPUT)
workbook = load_workbook(OUTPUT, keep_links=True)
sheet = workbook["模板"]

fields = {
    str(sheet.cell(5, col).value).strip(): col
    for col in range(1, sheet.max_column + 1)
    if sheet.cell(5, col).value
}
sku_col = fields["contribution_sku#1.value"]
title_col = fields["item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"]
description_col = fields["product_description[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"]
bullet_cols = [
    fields[f"bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#{index}.value"]
    for index in range(1, 6)
]

original_titles = {}
for row in range(7, 11):
    sku = str(sheet.cell(row, sku_col).value or "").strip()
    original_titles[sku] = sheet.cell(row, title_col).value
    bullets = list(COMMON_BULLETS)
    bullets[3] = DIMENSION_BULLETS[sku]
    listing = {"title": "", "description": DESCRIPTION}
    for index, bullet in enumerate(bullets, start=1):
        listing[f"bullet_{index}"] = bullet
    findings = validate_listing_row(listing)
    if findings:
        raise ValueError(f"Copy validation failed for {sku}: {findings}")
    sheet.cell(row, description_col).value = DESCRIPTION
    for col, bullet in zip(bullet_cols, bullets):
        sheet.cell(row, col).value = bullet

for row in range(7, 11):
    sku = str(sheet.cell(row, sku_col).value or "").strip()
    if sheet.cell(row, title_col).value != original_titles[sku]:
        raise ValueError(f"Title changed unexpectedly for {sku}")

workbook.calculation.fullCalcOnLoad = True
workbook.calculation.forceFullCalc = True
workbook.calculation.calcMode = "auto"
workbook.save(OUTPUT)
print(OUTPUT)
