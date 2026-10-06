from pathlib import Path

from openpyxl import load_workbook


OUT = (Path(__file__).resolve().parent / "../outputs").resolve()
SKU = "contribution_sku#1.value"
TITLE = "item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"


def open_sheet(name):
    path = OUT / name
    wb = load_workbook(path, keep_vba=True)
    ws = wb["Template"]
    fields = {str(ws.cell(5, col).value).strip(): col for col in range(1, ws.max_column + 1) if ws.cell(5, col).value}
    return path, wb, ws, fields


def fill_children(name, values):
    path, wb, ws, fields = open_sheet(name)
    for row in range(7, ws.max_row + 1):
        sku = str(ws.cell(row, fields[SKU]).value or "")
        if not sku or not any(sku.endswith(suffix) for suffix in values.get("suffixes", ())):
            continue
        for field_name, value in values["fields"].items():
            ws.cell(row, fields[field_name]).value = value
    wb.save(path)


apparel_fields = {
    "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.body_type": "Regular",
    "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.height_type": "Regular",
}
fill_children("内裤蝴蝶S码M码V1.xlsm", {"suffixes": ("-S", "-M"), "fields": {**apparel_fields, "special_size_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Standard"}})
fill_children("内裤棉M码5件装V1.xlsm", {"suffixes": ("-M-5pcs",), "fields": {**apparel_fields, "special_size_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Standard"}})
fill_children("袜子5双装V1.xlsm", {"suffixes": ("-5pairs",), "fields": apparel_fields})
fill_children("短袜5双装V1.xlsm", {"suffixes": ("-5pairs",), "fields": apparel_fields})
fill_children("船袜7双装V1.xlsm", {"suffixes": ("-7pairs",), "fields": apparel_fields})

for name, suffixes in [
    ("面罩迷彩全脸款V1.xlsm", ("-GreenCamo", "-BlueCamo")),
    ("面罩围脖款V1.xlsm", ("-DarkSet", "-LightSet")),
    ("冰袖冰川款V1.xlsm", ("-Glacier",)),
    ("冰袖渐变色V1.xlsm", ("-PinkYellow", "-BluePink", "-BlueYellow")),
]:
    fill_children(name, {"suffixes": suffixes, "fields": {"special_size_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Standard"}})

fill_children("骑行短裤M码2XL码V1.xlsm", {
    "suffixes": ("-M", "-2XL"),
    "fields": {
        "bottoms_size[marketplace_id=ATVPDKIKX0DER]#1.body_type": "Regular",
        "bottoms_size[marketplace_id=ATVPDKIKX0DER]#1.height_type": "Regular",
        "special_size_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Standard",
    },
})


title_updates = {
    "冰块托盘V1.xlsm": {
        "TTCA-IceTray-1pc": "Silicone Ice Cube Tray with Lid, Easy Release Round Ice Mold and Freezer Container for Cocktails and Drinks, Multicolor",
    },
    "骑行短裤M码2XL码V1.xlsm": {
        "TTCA-CyclingShorts": "5 Pack Men's Cycling Shorts, Quick Dry Athletic Bottoms for Running, Cycling, Gym Training, Multiple Styles Available",
        "TTCA-CyclingShorts-M": "5 Pack Men's Cycling Shorts, Quick Dry Athletic Bottoms for Running, Cycling and Gym Training, Multicolor, Size M",
        "TTCA-CyclingShorts-2XL": "5 Pack Men's Cycling Shorts, Quick Dry Athletic Bottoms for Running, Cycling and Gym Training, Multicolor, Size 2XL",
    },
    "充磁消磁器绿色V1.xlsm": {
        "TTCA-MagnetizerTool": "Screwdriver Magnetizer and Demagnetizer for Tips and Bits, Portable Accessory for Small Screws and Home Repair",
    },
    "直尺红色V1.xlsm": {
        "TTCA-MarkingRuler-Red": "Precision Marking Ruler with Sliding Stops, 11.81 Inch Aluminum Measuring Rule for Woodworking Layout, Red",
    },
    "拉力带V1.xlsm": {
        "TTCA-ExerciseBand": "Resistance Band for Pull Up Assistance, Stretching, Strength Training, Mobility and Fitness Workouts, Multicolor",
    },
    "方形烘焙纸50张V1.xlsm": {
        "TTCA-BakingPaper-50pcs": "50 Pack Square Parchment Paper Sheets, 8 by 6 Inches, Nonstick Baking Liners for Air Fryer, Oven and Cooking, Brown",
    },
}

for name, updates in title_updates.items():
    path, wb, ws, fields = open_sheet(name)
    for row in range(7, ws.max_row + 1):
        sku = str(ws.cell(row, fields[SKU]).value or "")
        if sku in updates:
            ws.cell(row, fields[TITLE]).value = updates[sku]
        if name == "方形烘焙纸50张V1.xlsm" and sku == "TTCA-BakingPaper-50pcs":
            size_field = "size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"
            ws.cell(row, fields[size_field]).value = "8 by 6 Inches"
    wb.save(path)


path, wb, ws, fields = open_sheet("搭扣20件装40毫米V1.xlsm")
fabric = "fabric_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"
for row in range(7, ws.max_row + 1):
    if ws.cell(row, fields[SKU]).value:
        ws.cell(row, fields[fabric]).value = "Metal"
wb.save(path)

print("final targeted fixes applied")
