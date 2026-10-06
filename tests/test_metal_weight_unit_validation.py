from openpyxl import Workbook

from app.template_validator import FIELD_NAMES, validate_template_file


def _make_template(path, metal_weight_unit):
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"
    headers = [
        FIELD_NAMES["sku"],
        "product_type#1.value",
        FIELD_NAMES["condition"],
        FIELD_NAMES["metal_weight_unit"],
    ]
    for column, field_name in enumerate(headers, 1):
        ws.cell(5, column).value = field_name
    ws.cell(7, 1).value = "TEST-SKU"
    ws.cell(7, 2).value = "APPAREL_PIN"
    ws.cell(7, 3).value = "New"
    ws.cell(7, 4).value = metal_weight_unit

    valid_values = wb.create_sheet("Valid Values")
    valid_values.cell(1, 1).value = "Metal Weight Unit - [ APPAREL_PIN ]"
    valid_values.cell(1, 2).value = "Grams"
    wb.save(path)


def test_metal_weight_unit_must_match_product_type_valid_values(tmp_path):
    path = tmp_path / "invalid.xlsx"
    _make_template(path, "Pounds")

    findings, _ = validate_template_file(path)

    assert any(item["field"] == "Metal Weight Unit" for item in findings)


def test_valid_metal_weight_unit_passes_enum_check(tmp_path):
    path = tmp_path / "valid.xlsx"
    _make_template(path, "Grams")

    findings, _ = validate_template_file(path)

    assert not any(item["field"] == "Metal Weight Unit" for item in findings)
