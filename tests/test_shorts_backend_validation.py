from openpyxl import Workbook

from app.template_validator import FIELD_NAMES, validate_template_file


def _make_template(path, *, fabric_type, bottoms_size):
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"
    headers = [
        FIELD_NAMES["sku"],
        "product_type#1.value",
        FIELD_NAMES["condition"],
        FIELD_NAMES["parentage_level"],
        FIELD_NAMES["fabric_type"],
        FIELD_NAMES["target_gender"],
        FIELD_NAMES["age_range"],
        FIELD_NAMES["bottoms_size_system"],
        FIELD_NAMES["bottoms_size_class"],
        FIELD_NAMES["bottoms_size"],
    ]
    for column, field_name in enumerate(headers, 1):
        ws.cell(5, column).value = field_name
    values = [
        "TEST-SHORTS-M",
        "SHORTS",
        "New",
        "Child",
        fabric_type,
        "Male",
        "Adult",
        "US",
        "Alpha",
        bottoms_size,
    ]
    for column, value in enumerate(values, 1):
        ws.cell(7, column).value = value
    wb.save(path)


def test_shorts_rejects_fabric_without_percentage_and_size_alias(tmp_path):
    path = tmp_path / "invalid.xlsx"
    _make_template(path, fabric_type="Polyester", bottoms_size="M")

    findings, _ = validate_template_file(path)

    assert any(item["field"] == "Fabric Type" for item in findings)
    assert any(item["field"] == "Bottoms Size Value" for item in findings)


def test_shorts_accepts_percentage_and_exact_size_enum(tmp_path):
    path = tmp_path / "valid.xlsx"
    _make_template(path, fabric_type="100% Polyester", bottoms_size="Medium")

    findings, _ = validate_template_file(path)

    assert not any(item["field"] in {"Fabric Type", "Bottoms Size Value"} for item in findings)
