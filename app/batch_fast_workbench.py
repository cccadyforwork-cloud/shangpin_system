import io
import json
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .paths import DATA_DIR, ROOT, safe_name


WORKBENCH_DIR = DATA_DIR / "batch_fast_workbench"
LEDGER_PATH = WORKBENCH_DIR / "ledger.json"
FILES_DIR = WORKBENCH_DIR / "files"
SOURCE_SHEETS_DIR = WORKBENCH_DIR / "source_sheets"

FILE_FIELDS = {"competitor_html", "source_template", "output_file"}
EDITABLE_FIELDS = FILE_FIELDS | {"status", "note"}
STATUS_VALUES = ("待处理", "已有模板", "生成中", "待复核", "可上传", "已上传", "需修正")
ALLOWED_SUFFIXES = {".html", ".htm", ".xlsx", ".xlsm", ".xls"}


def default_source_sheets():
    shared = [
        SOURCE_SHEETS_DIR / "2店批量上品.xlsx",
        SOURCE_SHEETS_DIR / "1店批量上品图图.xlsx",
        SOURCE_SHEETS_DIR / "1店批量上品凯旋.xlsx",
    ]
    if any(path.exists() for path in shared):
        return shared
    desktop = Path.home() / "Desktop"
    return [
        desktop / "2店批量上品" / "2店批量上品.xlsx",
        desktop / "1店批量上品图图" / "1店批量上品图图.xlsx",
        desktop / "1店批量上品凯旋" / "1店批量上品凯旋.xlsx",
    ]


def _resolve_stored_path(path_value):
    path = Path(str(path_value or "")).expanduser()
    return path.resolve() if path.is_absolute() else (ROOT / path).resolve()


def _store_path(path_value):
    path = _resolve_stored_path(path_value)
    try:
        return path.relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return str(path)


def load_ledger(path=LEDGER_PATH):
    path = Path(path)
    if not path.exists():
        return {"version": 1, "updated_at": "", "batches": []}
    data = json.loads(path.read_text(encoding="utf-8"))
    data.setdefault("version", 1)
    data.setdefault("updated_at", "")
    data.setdefault("batches", [])
    for batch in data["batches"]:
        for row in batch.get("rows", []):
            _normalize_output_versions(row)
    return data


def save_ledger(data, path=LEDGER_PATH):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    data["updated_at"] = datetime.now().astimezone().isoformat(timespec="seconds")
    temporary = path.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)
    return path


def ensure_ledger(path=LEDGER_PATH):
    data = load_ledger(path)
    if data["batches"]:
        return data
    sources = [source for source in default_source_sheets() if source.exists()]
    return import_sheets(sources, path) if sources else data


def import_sheets(paths=None, ledger_path=LEDGER_PATH):
    paths = list(paths or default_source_sheets())
    data = load_ledger(ledger_path)
    previous_by_source = {
        str(_resolve_stored_path(batch.get("source_sheet", ""))): batch
        for batch in data.get("batches", [])
        if batch.get("source_sheet")
    }
    batches = []
    imported = []
    seen_sources = set()
    for raw_path in paths:
        source = Path(raw_path).expanduser().resolve()
        source_key = str(source)
        if source_key in seen_sources or not source.exists() or source.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue
        batches.append(_batch_from_sheet(source, previous_by_source.get(source_key)))
        imported.append(source_key)
        seen_sources.add(source_key)
    for source_key, batch in previous_by_source.items():
        if source_key not in seen_sources:
            batches.append(batch)
    data["batches"] = batches
    data["last_imported"] = imported
    save_ledger(data, ledger_path)
    return data


def update_row(batch_id, row_id, fields, ledger_path=LEDGER_PATH):
    data = ensure_ledger(ledger_path)
    row = _find_row(data, batch_id, row_id)
    for key, value in (fields or {}).items():
        if key not in EDITABLE_FIELDS:
            continue
        text = str(value or "").strip()
        if key == "status" and text not in STATUS_VALUES:
            raise ValueError("状态值不正确。")
        if key == "output_file":
            if text:
                _register_output_version(row, text)
            else:
                row["output_file"] = ""
                row["output_versions"] = []
        else:
            row[key] = text
    save_ledger(data, ledger_path)
    return row


def attach_file(batch_id, row_id, field, filename, source, ledger_path=LEDGER_PATH, files_dir=FILES_DIR):
    if field not in {"source_template", "output_file"}:
        raise ValueError("这个位置不支持上传文件。")
    suffix = Path(filename or "").suffix.lower()
    if suffix not in {".xlsx", ".xlsm", ".xls"}:
        raise ValueError("模板和输出文档只支持 xlsx、xlsm、xls。")
    data = ensure_ledger(ledger_path)
    row = _find_row(data, batch_id, row_id)
    target_dir = Path(files_dir) / safe_name(batch_id) / safe_name(row_id) / safe_name(field)
    target_dir.mkdir(parents=True, exist_ok=True)
    target = _available_path(target_dir / safe_name(Path(filename).name))
    with target.open("wb") as output:
        shutil.copyfileobj(source, output)
    if field == "output_file":
        _register_output_version(row, target, filename=filename)
    else:
        row[field] = _store_path(target)
    if field == "source_template" and row.get("status") == "待处理":
        row["status"] = "已有模板"
    save_ledger(data, ledger_path)
    return target, row


def payload(ledger_path=LEDGER_PATH):
    data = ensure_ledger(ledger_path)
    batches = []
    totals = {"rows": 0, "html": 0, "templates": 0, "outputs": 0}
    for batch in data.get("batches", []):
        rendered_rows = []
        stores = batch.get("stores") or []
        primary_store = stores[0] if stores else ""
        for row in batch.get("rows", []):
            _normalize_output_versions(row)
            rendered = dict(row)
            rendered["files"] = {field: _file_info(row.get(field, "")) for field in FILE_FIELDS}
            versions = []
            current_path = str(row.get("output_file") or "")
            for item in sorted(row.get("output_versions", []), key=lambda entry: entry["version"], reverse=True):
                file_info = _file_info(item.get("path", ""))
                versions.append({
                    "version": item["version"],
                    "label": f"V{item['version']}",
                    "added_at": item.get("added_at", ""),
                    "current": str(item.get("path") or "") == current_path,
                    "file": file_info,
                })
            rendered["output_versions"] = versions
            rendered_rows.append(rendered)
            is_primary_product = not primary_store or row.get("store_id") == primary_store
            totals["rows"] += int(is_primary_product)
            totals["html"] += int(is_primary_product and rendered["files"]["competitor_html"]["exists"])
            totals["templates"] += int(rendered["files"]["source_template"]["exists"])
            totals["outputs"] += int(rendered["files"]["output_file"]["exists"])
        batches.append({**batch, "rows": rendered_rows})
    return {
        "ok": True,
        "updated_at": data.get("updated_at", ""),
        "statuses": list(STATUS_VALUES),
        "totals": totals,
        "batches": batches,
    }


def referenced_file(path_value, ledger_path=LEDGER_PATH, files_dir=FILES_DIR):
    raw = str(path_value or "").strip()
    if not raw:
        return None
    candidate = _resolve_stored_path(raw)
    data = ensure_ledger(ledger_path)
    allowed = set()
    for batch in data.get("batches", []):
        source_sheet = batch.get("source_sheet")
        if source_sheet:
            allowed.add(str(_resolve_stored_path(source_sheet)))
        for row in batch.get("rows", []):
            for field in FILE_FIELDS:
                if row.get(field):
                    allowed.add(str(_resolve_stored_path(row[field])))
            for item in row.get("output_versions", []):
                if item.get("path"):
                    allowed.add(str(_resolve_stored_path(item["path"])))
    return candidate if str(candidate) in allowed else None


def build_output_archive(batch_id, ledger_path=LEDGER_PATH, store_id=None):
    """Build an in-memory ZIP containing every available output in one batch."""
    data = ensure_ledger(ledger_path)
    batch = next((item for item in data.get("batches", []) if item.get("id") == batch_id), None)
    if batch is None:
        raise ValueError("找不到批次。")
    files = []
    for row in batch.get("rows", []):
        if store_id and row.get("store_id") != store_id:
            continue
        raw = str(row.get("output_file") or "").strip()
        path = _resolve_stored_path(raw) if raw else None
        if path and path.exists() and path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
            files.append(path)
    if not files:
        raise ValueError("当前批次还没有可下载的输出文档。")

    buffer = io.BytesIO()
    used_names = set()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            archive.write(path, arcname=_unique_archive_name(path.name, used_names))
    store_label = f"_{safe_name(store_id)}" if store_id else ""
    filename = f"{safe_name(batch.get('name') or batch_id)}{store_label}_输出文档.zip"
    return filename, buffer.getvalue(), len(files)


def package_shared_data(ledger_path=LEDGER_PATH, files_dir=FILES_DIR, source_sheets_dir=SOURCE_SHEETS_DIR):
    """Copy every ledger-referenced business file into the repository for Git sharing."""
    data = load_ledger(ledger_path)
    copied = 0
    missing = []
    for batch in data.get("batches", []):
        stores = batch.get("stores") or []
        primary_store = stores[0] if stores else ""
        shared_competitors = {}
        source_value = batch.get("source_sheet", "")
        if source_value:
            source = _resolve_stored_path(source_value)
            target = Path(source_sheets_dir) / safe_name(source.name)
            if source.is_file():
                if _path_is_within(source, source_sheets_dir):
                    batch["source_sheet"] = _store_path(source)
                else:
                    copied += _copy_shared_file(source, target)
                    batch["source_sheet"] = _store_path(target)
            else:
                missing.append(str(source))

        for row in batch.get("rows", []):
            for field in ("competitor_html", "source_template"):
                raw = str(row.get(field) or "").strip()
                if not raw:
                    continue
                product_id = str(row.get("asin") or row.get("id") or "")
                if field == "competitor_html" and primary_store and row.get("store_id") != primary_store:
                    shared_path = shared_competitors.get(product_id)
                    if shared_path:
                        row[field] = shared_path
                        continue
                source = _resolve_stored_path(raw)
                target = Path(files_dir) / safe_name(batch["id"]) / safe_name(row["id"]) / field / safe_name(source.name)
                if source.is_file():
                    if _path_is_within(source, files_dir):
                        row[field] = _store_path(source)
                    else:
                        copied += _copy_shared_file(source, target)
                        row[field] = _store_path(target)
                    if field == "competitor_html":
                        shared_competitors[product_id] = row[field]
                else:
                    missing.append(str(source))

            _normalize_output_versions(row)
            packaged_versions = []
            for item in row.get("output_versions", []):
                source = _resolve_stored_path(item["path"])
                target = Path(files_dir) / safe_name(batch["id"]) / safe_name(row["id"]) / "output_file" / safe_name(source.name)
                if source.is_file():
                    if _path_is_within(source, files_dir):
                        packaged_versions.append({**item, "path": _store_path(source)})
                    else:
                        copied += _copy_shared_file(source, target)
                        packaged_versions.append({**item, "path": _store_path(target)})
                else:
                    missing.append(str(source))
            row["output_versions"] = packaged_versions
            row["output_file"] = packaged_versions[-1]["path"] if packaged_versions else ""

    data["version"] = max(int(data.get("version") or 1), 2)
    save_ledger(data, ledger_path)
    return {"copied": copied, "missing": sorted(set(missing)), "ledger": str(Path(ledger_path).resolve())}


def _path_is_within(path, directory):
    try:
        Path(path).resolve().relative_to(Path(directory).resolve())
        return True
    except ValueError:
        return False


def _copy_shared_file(source, target):
    source = Path(source).resolve()
    target = Path(target).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    if source == target:
        return 0
    shutil.copy2(source, target)
    return 1


def _unique_archive_name(filename, used_names):
    candidate = filename
    index = 2
    path = Path(filename)
    while candidate.casefold() in used_names:
        candidate = f"{path.stem}_{index}{path.suffix}"
        index += 1
    used_names.add(candidate.casefold())
    return candidate


def _batch_from_sheet(path, previous=None):
    previous = previous or {}
    stores = [str(item).strip() for item in previous.get("stores", []) if str(item).strip()]
    primary_store = stores[0] if stores else ""
    previous_rows = {
        (str(row.get("store_id") or primary_store), str(row.get("asin") or row.get("id") or "")): row
        for row in previous.get("rows", [])
    }
    wb = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    ws = wb.active
    headers = {str(ws.cell(1, col).value or "").strip(): col for col in range(1, ws.max_column + 1)}
    link_col = headers.get("LINK", 1)
    html_col = headers.get("竞品HTML")
    template_col = headers.get("模版表格")
    note_col = headers.get("备注")
    rows = []
    for row_number in range(2, ws.max_row + 1):
        link = str(ws.cell(row_number, link_col).value or "").strip()
        if not link:
            continue
        match = re.search(r"/dp/([A-Z0-9]{10})", link, re.I)
        row_id = match.group(1).upper() if match else f"ROW-{row_number}"
        previous_row = previous_rows.get((primary_store, row_id), {})
        competitor_html = str(ws.cell(row_number, html_col).value or "").strip() if html_col else ""
        source_template = str(ws.cell(row_number, template_col).value or "").strip() if template_col else ""
        note = str(ws.cell(row_number, note_col).value or "").strip() if note_col else ""
        if (
            Path(competitor_html).suffix.lower() not in {".html", ".htm"}
            or not _resolve_stored_path(competitor_html).is_file()
        ):
            competitor_html = ""
        if (
            Path(source_template).suffix.lower() not in {".xlsx", ".xlsm", ".xls"}
            or not _resolve_stored_path(source_template).is_file()
        ):
            source_template = ""
        rendered_row = {
            "id": row_id,
            "asin": row_id if match else "",
            "row_number": row_number,
            "link": link,
            "competitor_html": competitor_html or _path_with_suffix(previous_row.get("competitor_html", ""), {".html", ".htm"}),
            "source_template": source_template or _path_with_suffix(previous_row.get("source_template", ""), {".xlsx", ".xlsm", ".xls"}),
            "output_file": previous_row.get("output_file", ""),
            "output_versions": previous_row.get("output_versions", []),
            "status": previous_row.get("status", "待处理"),
            "note": note or previous_row.get("note", ""),
        }
        if primary_store:
            rendered_row["store_id"] = primary_store
        rows.append(rendered_row)

    primary_rows = list(rows)
    for store_id in stores[1:]:
        for primary_row in primary_rows:
            product_id = str(primary_row.get("asin") or primary_row["id"])
            previous_row = previous_rows.get((store_id, product_id), {})
            rows.append({
                "id": previous_row.get("id") or _store_row_id(product_id, store_id),
                "asin": primary_row.get("asin", ""),
                "row_number": primary_row["row_number"],
                "link": primary_row["link"],
                "competitor_html": primary_row.get("competitor_html", ""),
                "source_template": _path_with_suffix(previous_row.get("source_template", ""), {".xlsx", ".xlsm", ".xls"}),
                "output_file": previous_row.get("output_file", ""),
                "output_versions": previous_row.get("output_versions", []),
                "status": previous_row.get("status", "待处理"),
                "note": previous_row.get("note", ""),
                "store_id": store_id,
            })

    batch = {
        "id": previous.get("id") or safe_name(path.stem).lower(),
        "name": previous.get("name") or path.stem,
        "source_sheet": _store_path(path),
        "rows": rows,
    }
    if stores:
        batch["stores"] = stores
    wb.close()
    return batch


def _store_row_id(product_id, store_id):
    return f"{product_id}--{safe_name(store_id)}"


def _find_row(data, batch_id, row_id):
    batch = next((item for item in data.get("batches", []) if item.get("id") == batch_id), None)
    if batch is None:
        raise ValueError("找不到批次。")
    row = next((item for item in batch.get("rows", []) if item.get("id") == row_id), None)
    if row is None:
        raise ValueError("找不到产品行。")
    return row


def _register_output_version(row, path_value, filename=None):
    _normalize_output_versions(row)
    path = _store_path(path_value)
    version = _output_version_number(filename or Path(path).name)
    existing = row.get("output_versions", [])
    if version is None:
        version = max((item["version"] for item in existing), default=0) + 1
    existing = [item for item in existing if item["version"] != version]
    existing.append({
        "version": version,
        "path": path,
        "added_at": datetime.now().astimezone().isoformat(timespec="seconds"),
    })
    row["output_versions"] = sorted(existing, key=lambda item: item["version"])
    row["output_file"] = max(row["output_versions"], key=lambda item: item["version"])["path"]


def _normalize_output_versions(row):
    normalized = {}
    for item in row.get("output_versions", []) or []:
        if not isinstance(item, dict) or not item.get("path"):
            continue
        try:
            version = int(item.get("version") or _output_version_number(Path(item["path"]).name) or 0)
        except (TypeError, ValueError):
            continue
        if version < 1:
            continue
        normalized[version] = {
            "version": version,
            "path": _store_path(item["path"]),
            "added_at": str(item.get("added_at") or ""),
        }
    current = str(row.get("output_file") or "").strip()
    if current:
        current_path = _store_path(current)
        current_version = _output_version_number(Path(current_path).name)
        if current_version is None:
            current_version = max(normalized, default=0) + 1
        normalized[current_version] = {
            "version": current_version,
            "path": current_path,
            "added_at": normalized.get(current_version, {}).get("added_at", ""),
        }
    row["output_versions"] = [normalized[key] for key in sorted(normalized)]
    row["output_file"] = row["output_versions"][-1]["path"] if row["output_versions"] else ""
    return row


def _output_version_number(filename):
    match = re.search(r"V(\d+)(?=\.[^.]+$)", str(filename or ""), re.I)
    return int(match.group(1)) if match else None


def _file_info(path_value):
    text = str(path_value or "").strip()
    path = _resolve_stored_path(text) if text else None
    exists = bool(path and path.exists() and path.is_file())
    return {
        "path": str(path.resolve()) if path else "",
        "name": path.name if path else "",
        "exists": exists,
        "size_bytes": path.stat().st_size if exists else 0,
    }


def _available_path(path):
    if not path.exists():
        return path
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return path.with_name(f"{path.stem}_{timestamp}{path.suffix}")


def _path_with_suffix(path_value, suffixes):
    text = str(path_value or "").strip()
    return text if Path(text).suffix.lower() in suffixes else ""
