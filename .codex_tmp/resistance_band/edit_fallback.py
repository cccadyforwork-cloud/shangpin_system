from copy import copy
from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook


SOURCE = Path("/Users/cc/Desktop/弹力带价格刷新.xlsx")
OUTPUT = Path("/Users/cc/Documents/GitHub/shangpin_system/outputs/01a09a0d-5e1d-7473-8611-f221c4c977bf/弹力带价格刷新V1.xlsx")


def copy_row(ws, source_row, target_row):
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        target.value = source.value
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        if source.hyperlink:
            target._hyperlink = copy(source.hyperlink)
        if source.comment:
            target.comment = copy(source.comment)
    if source_row in ws.row_dimensions:
        ws.row_dimensions[target_row] = copy(ws.row_dimensions[source_row])


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
copyfile(SOURCE, OUTPUT)
workbook = load_workbook(OUTPUT, keep_links=True)
sheet = workbook["模板"]

# Shift the three existing child rows down by one, bottom-up.
copy_row(sheet, 9, 10)
copy_row(sheet, 8, 9)
copy_row(sheet, 7, 8)

# Keep the existing data-row formatting while making row 7 the parent.
for col in range(1, sheet.max_column + 1):
    sheet.cell(7, col).value = None

parent_values = {
    1: "CA-LLD-Single",
    2: "EXERCISE_BAND",
    3: "Edit (Partial Update)",
    4: "Parent",
    6: "COLOR",
    7: "TPE Resistance Bands Flat Exercise Stretch Bands for Yoga Pilates Home Fitness Training",
    8: "Generic",
    9: "運動和戶外活動 > 運動與健身設備 > 力量訓練器材 > 彈力帶 (exercise-resistance-bands)",
}
for col, value in parent_values.items():
    sheet.cell(7, col).value = value

for row in range(8, 11):
    sheet.cell(row, 3).value = "Edit (Partial Update)"
    sheet.cell(row, 4).value = "Child"
    sheet.cell(row, 5).value = "CA-LLD-Single"
    sheet.cell(row, 6).value = "COLOR"

workbook.calculation.fullCalcOnLoad = True
workbook.calculation.forceFullCalc = True
workbook.calculation.calcMode = "auto"
workbook.save(OUTPUT)
print(OUTPUT)
