import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from PIL import Image

from app.refurb_batch_workbench import (
    BATCH_ID,
    BATCH_NAME,
    STORE_ID,
    attach_reference_screenshot,
    ensure_ledger,
    export_generated_images_to_desktop,
    import_sheets,
    payload,
)
from app.refurb_image_workflow import (
    AMAZON_UPLOAD_TARGET_BYTES,
    _extract_amazon_gallery_urls,
    prepare_amazon_upload_image,
)


class RefurbBatchWorkbenchTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.ledger = self.root / "ledger.json"
        self.sheet = self.root / "退货翻新批量上品.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["link", "品名", "竞品HTML", "模版表格", "备注"])
        sheet.append(["https://www.amazon.com/dp/B0ABCDEFGH", "隔油勺", "", "", "待翻新"])
        workbook.save(self.sheet)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_empty_workbench_has_only_cady_store_one(self):
        data = ensure_ledger(self.ledger)
        self.assertEqual(1, len(data["batches"]))
        batch = data["batches"][0]
        self.assertEqual(BATCH_ID, batch["id"])
        self.assertEqual(BATCH_NAME, batch["name"])
        self.assertEqual([STORE_ID], batch["stores"])

    def test_import_forces_every_row_to_cady_store_one(self):
        import_sheets([self.sheet], self.ledger)
        rendered = payload(self.ledger)
        self.assertEqual(1, len(rendered["batches"]))
        batch = rendered["batches"][0]
        self.assertEqual(BATCH_ID, batch["id"])
        self.assertEqual([STORE_ID], batch["stores"])
        self.assertEqual({STORE_ID}, {row["store_id"] for row in batch["rows"]})
        self.assertEqual("隔油勺", batch["rows"][0]["product_name"])
        self.assertEqual("待翻新", batch["rows"][0]["note"])

    def test_import_deduplicates_the_same_amazon_asin(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["link", "品名"])
        sheet.append(["https://www.amazon.com/dp/B0ABCDEFGH?ref=first", "镊子"])
        sheet.append(["https://www.amazon.com/dp/B0ABCDEFGH?ref=second", "镊子"])
        workbook.save(self.sheet)

        import_sheets([self.sheet], self.ledger)
        rows = payload(self.ledger)["batches"][0]["rows"]
        self.assertEqual(1, len(rows))
        self.assertEqual("https://www.amazon.com/dp/B0ABCDEFGH?ref=first", rows[0]["link"])

    def test_size_only_underwear_rows_are_merged_in_payload(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["link", "品名", "MSKU"])
        for row_number in range(2, 20):
            values = {
                12: ["https://a.co/d/yoga-xl", "内裤-瑜伽", "SP-underwear-yujia-XL"],
                14: ["https://a.co/d/yoga-l", "内裤-瑜伽", "SP-underwear-yujia-L"],
                15: ["https://a.co/d/lace-m", "内裤-蕾丝", "SP-underwear-leisi-M"],
                19: ["https://a.co/d/lace-s", "内裤-蕾丝", "SP-underwear-leisi-S"],
            }.get(row_number, ["", "", ""])
            sheet.append(values)
        workbook.save(self.sheet)

        import_sheets([self.sheet], self.ledger)
        rows = payload(self.ledger)["batches"][0]["rows"]

        self.assertEqual(2, len(rows))
        yoga = next(row for row in rows if row["product_name"] == "内裤-瑜伽")
        lace = next(row for row in rows if row["product_name"] == "内裤-蕾丝")
        self.assertEqual(["L", "XL"], yoga["variant_sizes"])
        self.assertEqual(["S", "M"], lace["variant_sizes"])
        self.assertEqual(2, len(yoga["source_links"]))
        self.assertEqual(2, len(lace["source_links"]))

    def test_confirmed_groups_merge_and_duplicate_tweezer_source_row_is_retained(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["link", "品名", "MSKU"])
        rows = {
            22: ["https://www.amazon.com/dp/B0BUTTER01", "内裤-蝴蝶", "SP-underwear-hudie-S"],
            25: ["https://www.amazon.com/dp/B0BUTTER02", "内裤-蝴蝶", "SP-underwear-hudie-M"],
            50: ["https://www.amazon.com/dp/B0SLEEVE01", "冰袖", "SP-BINGXIU-feng01"],
            51: ["https://www.amazon.com/dp/B0SLEEVE02", "冰袖", "SP-BINGXIU-lan01"],
            52: ["https://www.amazon.com/dp/B0SLEEVE03", "冰袖", "SP-BINGXIU-lan02"],
            53: ["https://www.amazon.com/dp/B0TWEEZER1", "镊子", "SP-NZ-X01"],
            54: ["https://www.amazon.com/dp/B0TWEEZER2", "镊子", "SP-NZ-W01"],
            55: ["https://www.amazon.com/dp/B0TWEEZER3", "镊子", "SP-NZ-J01"],
            56: ["https://www.amazon.com/dp/B0TWEEZER3", "镊子", "SP-NZ-Y01"],
        }
        for row_number in range(2, 57):
            sheet.append(rows.get(row_number, ["", "", ""]))
        workbook.save(self.sheet)

        import_sheets([self.sheet], self.ledger)
        rendered = payload(self.ledger)["batches"][0]["rows"]

        self.assertEqual(3, len(rendered))
        butterfly = next(row for row in rendered if row.get("merge_group") == "butterfly-underwear")
        sleeves = next(row for row in rendered if row.get("merge_group") == "gradient-arm-sleeves")
        tweezers = next(row for row in rendered if row.get("merge_group") == "tweezer-four-piece-set")
        self.assertEqual(["S", "M"], butterfly["variant_values"])
        self.assertEqual([50, 51, 52], sleeves["row_numbers"])
        self.assertEqual([53, 54, 55, 56], tweezers["row_numbers"])
        self.assertEqual(4, len(tweezers["source_links"]))
        self.assertEqual("bundle", tweezers["merge_type"])

    def test_variation_merge_exposes_independent_image_groups(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["link", "品名"])
        for row_number in range(2, 61):
            values = {
                5: ["https://www.amazon.com/dp/B0PHONE001", "手机架灰色"],
                60: ["https://www.amazon.com/dp/B0PHONE002", "手机架银色"],
            }.get(row_number, ["", ""])
            sheet.append(values)
        workbook.save(self.sheet)

        import_sheets([self.sheet], self.ledger)
        rendered = payload(self.ledger)["batches"][0]["rows"]
        phone_stand = next(row for row in rendered if row.get("merge_group") == "adjustable-phone-stand-colors")
        self.assertEqual(["Gray", "Silver"], [group["label"] for group in phone_stand["image_groups"]])
        self.assertEqual(2, len({group["key"] for group in phone_stand["image_groups"]}))

        files_dir = self.root / "files"
        gray_key = phone_stand["image_groups"][0]["key"]
        attach_reference_screenshot(
            BATCH_ID,
            phone_stand["id"],
            "gray.png",
            io.BytesIO(b"gray"),
            gray_key,
            self.ledger,
            files_dir,
        )
        refreshed = payload(self.ledger)["batches"][0]["rows"]
        phone_stand = next(row for row in refreshed if row.get("merge_group") == "adjustable-phone-stand-colors")
        counts = {group["label"]: group["reference_count"] for group in phone_stand["image_groups"]}
        self.assertEqual({"Gray": 1, "Silver": 0}, counts)

    def test_bundle_merge_keeps_one_combined_image_group(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["link", "品名"])
        for row_number in range(2, 57):
            values = {
                53: ["https://www.amazon.com/dp/B0TWEEZER1", "镊子X"],
                54: ["https://www.amazon.com/dp/B0TWEEZER2", "镊子W"],
                55: ["https://www.amazon.com/dp/B0TWEEZER3", "镊子J"],
                56: ["https://www.amazon.com/dp/B0TWEEZER3", "镊子Y"],
            }.get(row_number, ["", ""])
            sheet.append(values)
        workbook.save(self.sheet)

        import_sheets([self.sheet], self.ledger)
        rows = payload(self.ledger)["batches"][0]["rows"]
        tweezers = next(row for row in rows if row.get("merge_group") == "tweezer-four-piece-set")
        self.assertEqual(1, len(tweezers["image_groups"]))
        self.assertEqual("bundle", tweezers["image_groups"][0]["key"])
        self.assertEqual("组合装", tweezers["image_groups"][0]["label"])

    def test_extracts_amazon_gallery_images_in_display_order(self):
        page = """
        'colorImages': { 'initial': A.$.parseJSON('[{"hiRes":"https://m.media-amazon.com/images/I/main.jpg","large":"https://m.media-amazon.com/images/I/main-small.jpg"},{"hiRes":"https://m.media-amazon.com/images/I/detail.jpg"}]') }
        """
        self.assertEqual(
            [
                "https://m.media-amazon.com/images/I/main.jpg",
                "https://m.media-amazon.com/images/I/detail.jpg",
            ],
            _extract_amazon_gallery_urls(page),
        )

    def test_extracts_mobile_amazon_gallery_images(self):
        page = """
        <div id="image-block">
          <img data-a-hires="https://m.media-amazon.com/images/I/main._AC_UF894,1000_QL80_FMwebp_.jpg">
          <img data-a-hires="https://m.media-amazon.com/images/I/detail._AC_UF894,1000_QL80_FMwebp_.jpg">
        </div>
        """
        self.assertEqual(
            [
                "https://m.media-amazon.com/images/I/main._AC_UF894,1000_QL80_FMwebp_.jpg",
                "https://m.media-amazon.com/images/I/detail._AC_UF894,1000_QL80_FMwebp_.jpg",
            ],
            _extract_amazon_gallery_urls(page),
        )

    def test_exports_generated_images_to_desktop_without_zip(self):
        import_sheets([self.sheet], self.ledger)
        rendered = payload(self.ledger)["batches"][0]["rows"][0]
        generated_dir = self.root / "generated"
        generated_dir.mkdir()
        first = generated_dir / "隔油勺-1.png"
        second = generated_dir / "隔油勺-2.png"
        first.write_bytes(b"first")
        second.write_bytes(b"second")
        data = __import__("json").loads(self.ledger.read_text())
        row = data["batches"][0]["rows"][0]
        row["image_groups"] = {"default": {
            "generated_files": [str(first), str(second)],
            "status": "已生成",
        }}
        self.ledger.write_text(__import__("json").dumps(data, ensure_ascii=False))

        result = export_generated_images_to_desktop(
            BATCH_ID,
            rendered["id"],
            "default",
            self.ledger,
            self.root / "Desktop",
        )
        target = Path(result["path"])
        self.assertEqual("本轮第01批_进行中", target.parent.name)
        self.assertEqual(2, result["count"])
        self.assertEqual(["隔油勺-1.png", "隔油勺-2.png"], sorted(path.name for path in target.iterdir()))
        self.assertFalse(any(path.suffix == ".zip" for path in target.iterdir()))

    def test_prepares_generated_image_for_amazon_upload(self):
        source = self.root / "large.png"
        Image.effect_noise((2048, 2048), 100).convert("RGB").save(source)

        result = prepare_amazon_upload_image(source)

        self.assertEqual(".jpg", result.suffix)
        self.assertFalse(source.exists())
        self.assertLessEqual(result.stat().st_size, AMAZON_UPLOAD_TARGET_BYTES)
        with Image.open(result) as image:
            self.assertEqual("JPEG", image.format)
            self.assertEqual("RGB", image.mode)
            self.assertEqual((2048, 2048), image.size)
            self.assertGreaterEqual(image.info.get("dpi", (0, 0))[0], 72)


if __name__ == "__main__":
    unittest.main()
