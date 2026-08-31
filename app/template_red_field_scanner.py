"""Evaluate Amazon template conditional formatting without opening WPS/Excel.

Amazon category templates encode required and conditionally required cells as
formula-driven red borders.  This module evaluates the small formula subset used
by those rules and reports the cells that would render red for populated rows.
"""

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter


class FormulaNotSupported(ValueError):
    pass


@dataclass
class _Context:
    workbook: object
    sheet: object
    target_row: int
    target_col: int
    anchor_row: int
    anchor_col: int
    names: dict
    name_stack: tuple = ()


def scan_template_red_fields(path):
    """Return red-border findings and formulas that could not be evaluated."""
    path = Path(path)
    workbook = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    sheet = workbook["Template"] if "Template" in workbook.sheetnames else workbook.active
    names = {name.casefold(): item for name, item in workbook.defined_names.items()}
    data_rows = _data_rows(sheet)
    findings = []
    unresolved = []

    for conditional, rules in sheet.conditional_formatting._cf_rules.items():
        for cell_range in conditional.sqref.ranges:
            relevant_rows = [row for row in data_rows if cell_range.min_row <= row <= cell_range.max_row]
            if not relevant_rows:
                continue
            for col in range(cell_range.min_col, cell_range.max_col + 1):
                for row in relevant_rows:
                    context = _Context(
                        workbook=workbook,
                        sheet=sheet,
                        target_row=row,
                        target_col=col,
                        anchor_row=cell_range.min_row,
                        anchor_col=cell_range.min_col,
                        names=names,
                    )
                    matched = None
                    for rule in sorted(rules, key=lambda item: item.priority):
                        if rule.type != "expression" or not rule.formula:
                            continue
                        formula = rule.formula[0]
                        try:
                            applies = _truthy(_evaluate_formula(formula, context))
                        except FormulaNotSupported as exc:
                            unresolved.append({
                                "coordinate": f"{get_column_letter(col)}{row}",
                                "formula": formula,
                                "message": str(exc),
                            })
                            matched = "unresolved"
                            break
                        if applies:
                            matched = rule
                            if rule.stopIfTrue:
                                break
                    if matched in (None, "unresolved") or not _has_red_border(workbook, matched):
                        continue
                    sku = _row_sku(sheet, row)
                    findings.append({
                        "row": row,
                        "column": col,
                        "coordinate": f"{get_column_letter(col)}{row}",
                        "label": str(sheet.cell(4, col).value or "").strip(),
                        "field_name": str(sheet.cell(5, col).value or "").strip(),
                        "sku": sku,
                        "formula": matched.formula[0],
                        "kind": "required",
                    })
    workbook.close()
    return findings, _deduplicate_unresolved(unresolved)


def _data_rows(sheet):
    sku_col = None
    for col in range(1, sheet.max_column + 1):
        if str(sheet.cell(5, col).value or "").strip() == "contribution_sku#1.value":
            sku_col = col
            break
    if sku_col:
        return [row for row in range(7, sheet.max_row + 1) if sheet.cell(row, sku_col).value not in (None, "")]
    return [
        row for row in range(7, sheet.max_row + 1)
        if any(sheet.cell(row, col).value not in (None, "") for col in range(1, sheet.max_column + 1))
    ]


def _row_sku(sheet, row):
    for col in range(1, sheet.max_column + 1):
        if str(sheet.cell(5, col).value or "").strip() == "contribution_sku#1.value":
            return str(sheet.cell(row, col).value or "").strip()
    return ""


def _deduplicate_unresolved(items):
    result = []
    seen = set()
    for item in items:
        key = (item["coordinate"], item["formula"], item["message"])
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def _has_red_border(workbook, rule):
    differential = rule.dxf
    if differential is None and rule.dxfId is not None:
        differential = workbook._differential_styles.styles[rule.dxfId]
    border = getattr(differential, "border", None)
    if border is None:
        return False
    for side_name in ("left", "right", "top", "bottom"):
        side = getattr(border, side_name, None)
        color = getattr(side, "color", None)
        rgb = str(getattr(color, "rgb", "") or "").upper()
        indexed = getattr(color, "indexed", None)
        if rgb.endswith("FF0000") or indexed == 10:
            return True
    return False


def _evaluate_formula(formula, context):
    text = str(formula or "").strip()
    if not text.startswith("="):
        text = "=" + text
    tokens = [token for token in Tokenizer(text).items if token.type not in {"WSPACE", "WHITE-SPACE"}]
    parser = _FormulaParser(tokens, context)
    value = parser.parse_expression()
    if parser.position != len(tokens):
        raise FormulaNotSupported(f"无法解析公式尾部：{tokens[parser.position].value}")
    return value


class _FormulaParser:
    def __init__(self, tokens, context):
        self.tokens = tokens
        self.context = context
        self.position = 0

    def parse_expression(self):
        left = self._parse_primary()
        while self.position < len(self.tokens) and self.tokens[self.position].type == "OPERATOR-INFIX":
            operator = self.tokens[self.position].value
            if operator not in {"=", "<>", ">", "<", ">=", "<="}:
                raise FormulaNotSupported(f"暂不支持运算符：{operator}")
            self.position += 1
            right = self._parse_primary()
            left = _compare(left, right, operator)
        return left

    def _parse_primary(self):
        if self.position >= len(self.tokens):
            raise FormulaNotSupported("公式意外结束")
        token = self.tokens[self.position]
        self.position += 1
        if token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value[:-1].upper()
            arguments = []
            if not self._is_function_close():
                while True:
                    arguments.append(self.parse_expression())
                    if self._is_argument_separator():
                        self.position += 1
                        continue
                    break
            if not self._is_function_close():
                raise FormulaNotSupported(f"函数 {name} 缺少右括号")
            self.position += 1
            return _call_function(name, arguments)
        if token.type == "PAREN" and token.subtype == "OPEN":
            value = self.parse_expression()
            if self.position >= len(self.tokens) or self.tokens[self.position].subtype != "CLOSE":
                raise FormulaNotSupported("括号没有闭合")
            self.position += 1
            return value
        if token.type == "OPERAND":
            if token.subtype == "NUMBER":
                return float(token.value) if "." in token.value else int(token.value)
            if token.subtype == "TEXT":
                return token.value[1:-1].replace('""', '"')
            if token.subtype == "LOGICAL":
                return token.value.upper() == "TRUE"
            if token.subtype == "RANGE":
                return _resolve_reference(token.value, self.context)
        raise FormulaNotSupported(f"暂不支持公式标记：{token.value}")

    def _is_function_close(self):
        return (
            self.position < len(self.tokens)
            and self.tokens[self.position].type == "FUNC"
            and self.tokens[self.position].subtype == "CLOSE"
        )

    def _is_argument_separator(self):
        return self.position < len(self.tokens) and self.tokens[self.position].type == "SEP"


def _call_function(name, arguments):
    if name == "AND":
        return all(_truthy(value) for value in arguments)
    if name == "OR":
        return any(_truthy(value) for value in arguments)
    if name == "NOT" and len(arguments) == 1:
        return not _truthy(arguments[0])
    if name == "IF" and len(arguments) in (2, 3):
        return arguments[1] if _truthy(arguments[0]) else (arguments[2] if len(arguments) == 3 else False)
    if name == "LEN" and len(arguments) == 1:
        value = arguments[0]
        return len(str(value)) if value not in (None, "") else 0
    if name == "COUNTIF" and len(arguments) == 2:
        values = arguments[0] if isinstance(arguments[0], list) else [arguments[0]]
        criterion = arguments[1]
        return sum(1 for value in values if _excel_equal(value, criterion))
    if name == "ISNUMBER" and len(arguments) == 1:
        return isinstance(arguments[0], (int, float)) and not isinstance(arguments[0], bool)
    raise FormulaNotSupported(f"暂不支持函数或参数数量：{name}/{len(arguments)}")


def _resolve_reference(reference, context):
    reference = reference.strip()
    direct = _parse_direct_reference(reference)
    if direct:
        return _read_direct_reference(direct, context)
    name_key = reference.casefold()
    defined = context.names.get(name_key)
    if defined is None:
        raise FormulaNotSupported(f"找不到名称：{reference}")
    if name_key in context.name_stack:
        raise FormulaNotSupported(f"名称循环引用：{reference}")
    attr_text = str(defined.attr_text or "").strip()
    if attr_text.startswith("="):
        attr_text = attr_text[1:]
    direct = _parse_direct_reference(attr_text)
    nested = _Context(
        workbook=context.workbook,
        sheet=context.sheet,
        target_row=context.target_row,
        target_col=context.target_col,
        anchor_row=1,
        anchor_col=1,
        names=context.names,
        name_stack=context.name_stack + (name_key,),
    )
    if direct:
        return _read_direct_reference(direct, nested)
    return _evaluate_formula(attr_text, nested)


_CELL_RANGE_RE = re.compile(
    r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?"
    r"(?P<cabs1>\$?)(?P<col1>[A-Z]{1,3})(?P<rabs1>\$?)(?P<row1>\d+)"
    r"(?::(?P<cabs2>\$?)(?P<col2>[A-Z]{1,3})(?P<rabs2>\$?)(?P<row2>\d+))?$",
    re.I,
)


def _parse_direct_reference(value):
    return _CELL_RANGE_RE.match(value.strip())


def _read_direct_reference(match, context):
    sheet_name = match.group("sheet")
    if sheet_name:
        if sheet_name.startswith("'"):
            sheet_name = sheet_name[1:-1].replace("''", "'")
        sheet = context.workbook[sheet_name]
    else:
        sheet = context.sheet

    def adjusted(col_text, row_text, col_absolute, row_absolute):
        col = column_index_from_string(col_text)
        row = int(row_text)
        if not col_absolute:
            col += context.target_col - context.anchor_col
        if not row_absolute:
            row += context.target_row - context.anchor_row
        return row, col

    start = adjusted(match.group("col1"), match.group("row1"), match.group("cabs1"), match.group("rabs1"))
    if not match.group("col2"):
        return sheet.cell(*start).value
    end = adjusted(match.group("col2"), match.group("row2"), match.group("cabs2"), match.group("rabs2"))
    min_row, max_row = sorted((start[0], end[0]))
    min_col, max_col = sorted((start[1], end[1]))
    return [
        sheet.cell(row, col).value
        for row in range(min_row, max_row + 1)
        for col in range(min_col, max_col + 1)
    ]


def _truthy(value):
    if isinstance(value, list):
        return bool(value)
    if value in (None, "", 0, 0.0, False):
        return False
    return True


def _excel_equal(left, right):
    if left in (None, "") and right in (None, ""):
        return True
    if isinstance(left, str) or isinstance(right, str):
        return str(left or "").casefold() == str(right or "").casefold()
    return left == right


def _compare(left, right, operator):
    if operator == "=":
        return _excel_equal(left, right)
    if operator == "<>":
        return not _excel_equal(left, right)
    try:
        if operator == ">":
            return left > right
        if operator == "<":
            return left < right
        if operator == ">=":
            return left >= right
        if operator == "<=":
            return left <= right
    except TypeError:
        left, right = str(left or ""), str(right or "")
        return {">": left > right, "<": left < right, ">=": left >= right, "<=": left <= right}[operator]
    raise FormulaNotSupported(f"暂不支持比较符：{operator}")
