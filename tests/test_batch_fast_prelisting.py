import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook

from app.batch_fast_prelisting import (
    _child_sku,
    _fast_copy_fallback,
    _output_filename,
    _parent_required_overlay_fields,
    _parent_sku,
    _resolve_task_price,
    _row_changed,
    extract_competitor_price,
    extract_competitor_title,
    extract_competitor_variants,
    resolve_logistics_tier,
    rewrite_competitor_title_tail,
)
from app.listing_rules import validate_listing_row


class BatchFastPrelistingTest(unittest.TestCase):
    def test_resolves_le_4_oz_from_page_weight(self):
        tier = resolve_logistics_tier(weight_grams=80)
        self.assertEqual(tier["id"], "le_4_oz")
        self.assertEqual(tier["package_in"], [3.94, 3.94, 0.79])
        self.assertEqual(tier["weight_lb"], 0.22)

    def test_resolves_explicit_tier(self):
        tier = resolve_logistics_tier(tier_id="4_8_oz")
        self.assertEqual(tier["package_in"], [5.91, 3.94, 1.57])
        self.assertEqual(tier["weight_lb"], 0.44)

    def test_pet_toy_parent_gets_item_measurements_only(self):
        tier = resolve_logistics_tier(tier_id="le_4_oz")
        fields = _parent_required_overlay_fields("PET_TOY", tier)
        self.assertEqual(fields["item_length_width_height[marketplace_id=ATVPDKIKX0DER]#1.length.value"], 3.94)
        self.assertEqual(fields["item_weight[marketplace_id=ATVPDKIKX0DER]#1.value"], 0.22)
        self.assertFalse(any("item_package_dimensions" in field for field in fields))
        self.assertFalse(any("item_package_weight" in field for field in fields))
        self.assertEqual(_parent_required_overlay_fields("HAT", tier), {})

    def test_towel_and_gift_wrap_parent_backend_requirements(self):
        tier = resolve_logistics_tier(tier_id="4_8_oz")
        towel = _parent_required_overlay_fields("TOWEL", tier, {"size": "9.84 By 9.84 Inches"})
        self.assertEqual(towel["item_length_width[marketplace_id=ATVPDKIKX0DER]#1.length.value"], 9.84)
        self.assertEqual(towel["item_length_width[marketplace_id=ATVPDKIKX0DER]#1.width.value"], 9.84)
        self.assertEqual(towel["item_weight[marketplace_id=ATVPDKIKX0DER]#1.value"], 0.44)
        gift_wrap = _parent_required_overlay_fields("GIFT_WRAP", tier, {"set_count": 50})
        self.assertEqual(gift_wrap["unit_count[marketplace_id=ATVPDKIKX0DER]#1.value"], 50)
        self.assertEqual(gift_wrap["item_weight[marketplace_id=ATVPDKIKX0DER]#1.value"], 0.44)

    def test_short_sku_generation(self):
        task = {"product_name": "glass BEADS extra words", "color": "Green", "variation_theme": "COLOR"}
        self.assertEqual(_parent_sku(task), "CA-GlassBeads")
        self.assertEqual(_child_sku(task, {}, 1), "CA-GlassBeads-Green")

    def test_number_of_items_sku_uses_lowercase_unit_suffix(self):
        task = {"product_name": "tool SHARPENER", "set_count": 2, "variation_theme": "NUMBER_OF_ITEMS"}
        self.assertEqual(_child_sku(task, {}, 1), "CA-ToolSharpener-2pcs")

    def test_size_sku_uses_lowercase_measurement_unit(self):
        task = {
            "product_name": "storage hook",
            "size": "6 MM",
            "variation_theme": "SIZE",
        }

        self.assertEqual(_child_sku(task, {}, 1), "CA-StorageHook-6mm")

    def test_non_unit_variant_sku_remains_title_case(self):
        task = {
            "product_name": "glass beads",
            "color": "Dark Green",
            "variation_theme": "COLOR",
        }

        self.assertEqual(_child_sku(task, {}, 1), "CA-GlassBeads-DarkGreen")

    def test_explicit_sku_remains_an_override(self):
        task = {"product_name": "glass beads", "parent_sku": "CUSTOM-P", "child_sku": "CUSTOM-C"}
        self.assertEqual(_parent_sku(task), "CUSTOM-P")
        self.assertEqual(_child_sku(task, {}, 1), "CUSTOM-C")

    def test_output_name_cannot_escape_output_directory(self):
        filename = _output_filename({"output_name": "../危险文件.xlsx"}, "fallback", ".xlsm")
        self.assertEqual(filename, "危险文件V1.xlsx")

    def test_fallback_copy_passes_listing_rules(self):
        copy = _fast_copy_fallback({
            "product_name": "Glass Beads",
            "color": "Green",
            "size": "6 MM",
            "material": "Glass",
            "set_count": 100,
        })
        self.assertEqual(len(copy["description"].split("\n\n")), 4)
        self.assertGreaterEqual(len(copy["description"]), 1500)
        self.assertLessEqual(len(copy["description"]), 1800)
        self.assertEqual(validate_listing_row(copy), [])

    def test_extracts_complete_single_color_twister_without_asins(self):
        html = '''
        <input name="twisterDimKeys" value="color_name"/>
        <li id="color_name_0" title="Click to select Blue" data-csa-c-item-id="ASIN1"></li>
        <li id="color_name_1" title="Click to select Llight Blue" data-csa-c-item-id="ASIN2"></li>
        <li id="color_name_2" title="Click to select Green" data-csa-c-item-id="ASIN3"></li>
        '''
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "competitor.html"
            path.write_text(html, encoding="utf-8")
            variants = extract_competitor_variants([path])
        self.assertEqual(variants, [
            {"color": "Blue"},
            {"color": "Light Blue"},
            {"color": "Green"},
        ])
        self.assertTrue(all("asin" not in variant for variant in variants))

    def test_does_not_auto_expand_combined_twister(self):
        html = '''
        <input name="twisterDimKeys" value="color_name,size_name"/>
        <li id="color_name_0" title="Click to select Blue"></li>
        '''
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "competitor.html"
            path.write_text(html, encoding="utf-8")
            self.assertEqual(extract_competitor_variants([path]), [])

    def test_extracts_current_price_not_discount_amount(self):
        html = '''
        <span class="aok-offscreen"> $1.95 </span>
        <span class="a-price priceToPay"><span class="a-offscreen"></span></span>
        <span>Includes $1.96 Amazon discount</span>
        '''
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "competitor.html"
            path.write_text(html, encoding="utf-8")
            self.assertEqual(extract_competitor_price([path]), 1.95)

    def test_price_resolution_prefers_saved_html_before_online_and_estimate(self):
        html = '''
        <span class="aok-offscreen">$1.95</span>
        <span class="a-price priceToPay"><span class="a-offscreen"></span></span>
        '''
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "competitor.html"
            path.write_text(html, encoding="utf-8")
            price = _resolve_task_price(
                {"online_price": 2.62, "estimated_price": 2.99},
                [path],
            )
            self.assertEqual(price, 1.95)

    def test_price_resolution_uses_online_then_estimated_price(self):
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "competitor.html"
            path.write_text("<html><body>Currently unavailable.</body></html>", encoding="utf-8")
            self.assertEqual(_resolve_task_price({"online_price": 2.62, "estimated_price": 2.99}, [path]), 2.62)
            self.assertEqual(_resolve_task_price({"estimated_price": 2.99}, [path]), 2.99)

    def test_extracts_amazon_product_title_before_page_title(self):
        html = '''
        <title>Fallback Title : Amazon.com: Home</title>
        <span id="productTitle">Original Product Keywords for Home Craft Projects and Daily Organization</span>
        '''
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "competitor.html"
            path.write_text(html, encoding="utf-8")
            self.assertEqual(
                extract_competitor_title([path]),
                "Original Product Keywords for Home Craft Projects and Daily Organization",
            )

    def test_competitor_title_adds_only_three_to_five_tail_words(self):
        source = "Compact Storage Hooks for Closet Shelves Travel Bags and Everyday Home Organization"
        rewritten = rewrite_competitor_title_tail(source)
        self.assertTrue(rewritten.startswith(source))
        added_words = rewritten[len(source):].strip(" ,").split()
        self.assertGreaterEqual(len(added_words), 3)
        self.assertLessEqual(len(added_words), 5)

    def test_long_competitor_title_deletes_only_three_to_five_tail_words(self):
        source = " ".join(f"Keyword{index}" for index in range(1, 19))
        rewritten = rewrite_competitor_title_tail(source)
        source_words = source.split()
        rewritten_words = rewritten.split()
        self.assertGreaterEqual(len(source_words) - len(rewritten_words), 3)
        self.assertLessEqual(len(source_words) - len(rewritten_words), 5)
        self.assertEqual(rewritten_words[:10], source_words[:10])
        source_iter = iter(source_words)
        self.assertTrue(all(any(candidate == word for candidate in source_iter) for word in rewritten_words))
        self.assertGreaterEqual(len(rewritten), 100)
        self.assertLessEqual(len(rewritten), 125)

    def test_row_change_detection_ignores_action_only(self):
        ws = Workbook().active
        ws.cell(1, 1).value = "Same Title"
        ws.cell(1, 2).value = "Edit (Partial Update)"
        fields = {"item_name": 1, "::record_action": 2}
        previous = {"item_name": "Same Title"}
        self.assertFalse(_row_changed(ws, 1, fields, previous))
        ws.cell(1, 1).value = "Updated Title"
        self.assertTrue(_row_changed(ws, 1, fields, previous))


if __name__ == "__main__":
    unittest.main()
