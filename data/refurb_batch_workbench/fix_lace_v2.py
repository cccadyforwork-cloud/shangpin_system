from pathlib import Path
import shutil

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/refurb_batch_workbench/files/cady退货翻新/ROW-15/output_file/内裤蕾丝V1.xlsm"
OUTPUT = ROOT / "data/refurb_batch_workbench/files/cady退货翻新/ROW-15/output_file/内裤蕾丝V2.xlsm"


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    wb = load_workbook(OUTPUT, keep_vba=True)
    ws = wb["Template"]
    fields = {
        str(ws.cell(5, col).value).strip(): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(5, col).value
    }
    sku_field = "contribution_sku#1.value"
    material_field = "material[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"
    fabric_field = "fabric_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"
    size_field = "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size"
    size_to_field = "apparel_size[marketplace_id=ATVPDKIKX0DER]#1.size_to"
    size_by_sku = {
        "TTCA-LaceUnderwear-S": "Small",
        "TTCA-LaceUnderwear-M": "Medium",
    }

    for row in range(7, ws.max_row + 1):
        sku = ws.cell(row, fields[sku_field]).value
        if not sku:
            continue
        ws.cell(row, fields[material_field]).value = "Viscose and Spandex"
        ws.cell(row, fields[fabric_field]).value = "95% Viscose, 5% Spandex"
        if sku in size_by_sku:
            ws.cell(row, fields[size_field]).value = size_by_sku[sku]
            ws.cell(row, fields[size_to_field]).value = None

    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
