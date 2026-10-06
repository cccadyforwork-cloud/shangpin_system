import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from .batch_fast_workbench import (
    attach_file as _attach_file,
    build_output_archive as _build_output_archive,
    import_sheets as _import_sheets,
    load_ledger,
    package_shared_data as _package_shared_data,
    payload as _payload,
    referenced_file as _referenced_file,
    save_ledger,
    update_row as _update_row,
)
from .paths import DATA_DIR, safe_name
from .refurb_image_workflow import (
    attach_reference_screenshot as _attach_reference_screenshot,
    enrich_payload as _enrich_image_payload,
    export_generated_images_to_desktop as _export_generated_images_to_desktop,
    fetch_reference_images as _fetch_reference_images,
    referenced_image_file as _referenced_image_file,
    start_generation as _start_image_generation,
)


WORKBENCH_DIR = DATA_DIR / "refurb_batch_workbench"
LEDGER_PATH = WORKBENCH_DIR / "ledger.json"
FILES_DIR = WORKBENCH_DIR / "files"
SOURCE_SHEETS_DIR = WORKBENCH_DIR / "source_sheets"
BATCH_ID = "cady退货翻新"
BATCH_NAME = "Cady 退货翻新"
STORE_ID = "1店"
SIZE_ORDER = {"XS": 0, "S": 1, "M": 2, "L": 3, "XL": 4, "XXL": 5, "XXXL": 6}
MERGE_GROUPS = (
    {
        "id": "adjustable-phone-stand-colors",
        "rows": (5, 60),
        "product_name": "手机架",
        "variation_theme": "COLOR",
        "attribute_label": "颜色",
        "values": {5: "Gray", 60: "Silver"},
    },
    {
        "id": "mens-cycling-shorts-sizes",
        "rows": (72, 73),
        "product_name": "骑行短裤",
        "variation_theme": "SIZE",
        "attribute_label": "尺码",
        "values": {72: "M", 73: "2XL"},
    },
    {
        "id": "waterproof-phone-pouch-colors",
        "rows": (110, 111),
        "product_name": "手机防水袋",
        "variation_theme": "COLOR",
        "attribute_label": "颜色",
        "values": {110: "White", 111: "Black"},
    },
    {
        "id": "yoga-underwear",
        "rows": (14, 12),
        "product_name": "内裤-瑜伽",
        "variation_theme": "SIZE",
        "attribute_label": "尺码",
        "values": {12: "XL", 14: "L"},
    },
    {
        "id": "lace-underwear",
        "rows": (19, 15),
        "product_name": "内裤-蕾丝",
        "variation_theme": "SIZE",
        "attribute_label": "尺码",
        "values": {15: "M", 19: "S"},
    },
    {
        "id": "butterfly-underwear",
        "rows": (22, 25),
        "product_name": "内裤-蝴蝶",
        "variation_theme": "SIZE",
        "attribute_label": "尺码",
        "values": {22: "S", 25: "M"},
    },
    {
        "id": "gradient-arm-sleeves",
        "rows": (50, 51, 52),
        "product_name": "冰袖-渐变色",
        "variation_theme": "COLOR",
        "attribute_label": "颜色组合",
        "values": {50: "粉黄渐变", 51: "蓝粉渐变", 52: "蓝黄渐变"},
    },
    {
        "id": "solid-full-face-mask",
        "rows": (11, 36),
        "product_name": "面罩-纯色全脸款",
        "variation_theme": "COLOR",
        "attribute_label": "颜色组合",
        "values": {11: "深色组", 36: "亮色组"},
    },
    {
        "id": "camo-full-face-mask",
        "rows": (37, 39),
        "product_name": "面罩-迷彩全脸款",
        "variation_theme": "COLOR",
        "attribute_label": "颜色组合",
        "values": {37: "绿色迷彩", 39: "蓝色迷彩"},
    },
    {
        "id": "neck-gaiter-mask",
        "rows": (38, 40),
        "product_name": "面罩-围脖款",
        "variation_theme": "COLOR",
        "attribute_label": "颜色组合",
        "values": {38: "深色组", 40: "浅色组"},
    },
    {
        "id": "tweezer-four-piece-set",
        "rows": (53, 54, 55, 56),
        "product_name": "镊子-4件套",
        "variation_theme": "",
        "attribute_label": "组合",
        "values": {53: "X款", 54: "W款", 55: "J款", 56: "Y款"},
        "merge_type": "bundle",
    },
    {
        "id": "metal-hair-tie-style-color",
        "rows": (82, 83, 84, 85),
        "product_name": "蝴蝶结",
        "variation_theme": "STYLE/COLOR",
        "attribute_label": "款式 / 颜色",
        "values": {
            82: "双弧几何款 / 金色",
            83: "蝴蝶结款 / 银色",
            84: "蝴蝶结款 / 金色",
            85: "双弧几何款 / 银色",
        },
    },
    {
        "id": "dog-rope-toy-set",
        "rows": (112, 113, 114),
        "product_name": "狗玩具",
        "variation_theme": "SET_NAME",
        "attribute_label": "组合款式",
        "values": {
            112: "2件套球形组合",
            113: "3件套混合组合",
            114: "1件装长条绳结",
        },
    },
)
MERGE_GROUP_BY_ROW = {
    row_number: group
    for group in MERGE_GROUPS
    for row_number in group["rows"]
}
RETAIN_DUPLICATE_SOURCE_ROWS = {56}


def default_source_sheets():
    shared = SOURCE_SHEETS_DIR / "退货翻新批量上品.xlsx"
    if shared.exists():
        return [shared]
    shared_candidates = sorted(SOURCE_SHEETS_DIR.glob("*.xlsx"))
    if shared_candidates:
        return [shared_candidates[0]]
    desktop = Path.home() / "Desktop"
    return [
        desktop / "副本入库记录-退货入库-总和.xlsx",
        desktop / "退货翻新批量上品" / "退货翻新批量上品.xlsx",
    ]


def ensure_ledger(path=LEDGER_PATH):
    data = load_ledger(path)
    if data.get("batches"):
        return _normalize_workbench(data, path)
    sources = [source for source in default_source_sheets() if source.exists()]
    if sources:
        return import_sheets(sources, path)
    data["batches"] = [{
        "id": BATCH_ID,
        "name": BATCH_NAME,
        "source_sheet": "",
        "stores": [STORE_ID],
        "rows": [],
    }]
    save_ledger(data, path)
    return data


def import_sheets(paths=None, ledger_path=LEDGER_PATH):
    selected = list(paths or default_source_sheets())
    existing = [path for path in selected if Path(path).expanduser().exists()]
    if not existing:
        return ensure_ledger(ledger_path)
    data = _import_sheets(existing, ledger_path)
    return _normalize_workbench(data, ledger_path)


def payload(ledger_path=LEDGER_PATH):
    ensure_ledger(ledger_path)
    rendered = _payload(ledger_path)
    for batch in rendered.get("batches", []):
        batch["rows"] = _merge_confirmed_rows(batch.get("rows", []))
    rows = [row for batch in rendered.get("batches", []) for row in batch.get("rows", [])]
    rendered["totals"] = {
        "rows": len(rows),
        "html": sum(int(row["files"]["competitor_html"]["exists"]) for row in rows),
        "templates": sum(
            int(bool(row.get("source_template_files")) or row["files"]["source_template"]["exists"])
            for row in rows
        ),
        "outputs": sum(int(row["files"]["output_file"]["exists"]) for row in rows),
    }
    return _enrich_image_payload(rendered, ledger_path)


def update_row(batch_id, row_id, fields, ledger_path=LEDGER_PATH):
    ensure_ledger(ledger_path)
    return _update_row(batch_id, row_id, fields, ledger_path)


def attach_file(batch_id, row_id, field, filename, source, ledger_path=LEDGER_PATH, files_dir=FILES_DIR):
    ensure_ledger(ledger_path)
    return _attach_file(batch_id, row_id, field, filename, source, ledger_path, files_dir)


def referenced_file(path_value, ledger_path=LEDGER_PATH, files_dir=FILES_DIR):
    ensure_ledger(ledger_path)
    return _referenced_file(path_value, ledger_path, files_dir) or _referenced_image_file(path_value, ledger_path)


def attach_reference_screenshot(
    batch_id,
    row_id,
    filename,
    source,
    image_group_key="default",
    ledger_path=LEDGER_PATH,
    files_dir=FILES_DIR,
):
    ensure_ledger(ledger_path)
    primary_id, key, label, _product_name = _resolve_image_group(
        batch_id, row_id, image_group_key, ledger_path
    )
    return _attach_reference_screenshot(
        batch_id,
        primary_id,
        filename,
        source,
        ledger_path,
        files_dir,
        image_group_key=key,
        group_label=label,
    )


def start_image_generation(
    batch_id,
    row_id,
    image_group_key="default",
    ledger_path=LEDGER_PATH,
    files_dir=FILES_DIR,
):
    ensure_ledger(ledger_path)
    primary_id, key, label, product_name = _resolve_image_group(
        batch_id, row_id, image_group_key, ledger_path
    )
    return _start_image_generation(
        batch_id,
        primary_id,
        ledger_path,
        files_dir,
        image_group_key=key,
        group_label=label,
        product_name_override=product_name,
    )


def fetch_reference_images(
    batch_id,
    row_id,
    image_group_key="default",
    ledger_path=LEDGER_PATH,
    files_dir=FILES_DIR,
):
    ensure_ledger(ledger_path)
    primary_id, key, label, product_name = _resolve_image_group(
        batch_id, row_id, image_group_key, ledger_path
    )
    data = load_ledger(ledger_path)
    batch = next((item for item in data.get("batches", []) if item.get("id") == batch_id), None)
    rows = batch.get("rows", []) if batch else []
    requested = next((item for item in rows if item.get("id") == row_id), None)
    if requested is None:
        raise ValueError("找不到产品行。")
    merge_group = MERGE_GROUP_BY_ROW.get(requested.get("row_number"))
    if merge_group is None:
        source_urls = [requested.get("link", "")]
    else:
        members = [item for item in rows if item.get("row_number") in merge_group["rows"]]
        if merge_group.get("merge_type", "variation") == "bundle":
            source_urls = [item.get("link", "") for item in members]
        else:
            member = next((item for item in members if item.get("id") == key), None)
            source_urls = [member.get("link", "") if member else ""]
    return _fetch_reference_images(
        batch_id,
        primary_id,
        source_urls,
        ledger_path,
        files_dir,
        image_group_key=key,
        group_label=label,
        product_name_override=product_name,
    )


def export_generated_images_to_desktop(
    batch_id,
    row_id,
    image_group_key="default",
    ledger_path=LEDGER_PATH,
    desktop_dir=None,
):
    ensure_ledger(ledger_path)
    primary_id, key, label, product_name = _resolve_image_group(
        batch_id, row_id, image_group_key, ledger_path
    )
    return _export_generated_images_to_desktop(
        batch_id,
        primary_id,
        ledger_path,
        image_group_key=key,
        group_label=label,
        product_name_override=product_name,
        desktop_dir=desktop_dir,
    )


def build_all_image_groups_archive(batch_id, row_id, ledger_path=LEDGER_PATH):
    data = ensure_ledger(ledger_path)
    batch = next((item for item in data.get("batches", []) if item.get("id") == batch_id), None)
    if batch is None:
        raise ValueError("找不到批次。")
    row = next((item for item in batch.get("rows", []) if item.get("id") == row_id), None)
    if row is None:
        raise ValueError("找不到产品行。")
    group = MERGE_GROUP_BY_ROW.get(row.get("row_number"))
    if group is None or group.get("merge_type", "variation") == "bundle":
        raise ValueError("这个商品没有多套图组。")
    members = [item for item in batch.get("rows", []) if item.get("row_number") in group["rows"]]
    if not members or members[0].get("id") != row_id:
        raise ValueError("请从合并商品主行下载全部图组。")
    stored = row.get("image_groups", {}) or {}
    output = io.BytesIO()
    count = 0
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for member in members:
            key = member.get("id", "")
            label = group.get("values", {}).get(member.get("row_number"), "") or key
            state = stored.get(key, {})
            zip_value = state.get("zip", "")
            if not zip_value and member is members[0]:
                zip_value = row.get("image_zip", "")
            zip_path = Path(str(zip_value or ""))
            if not zip_path.is_absolute():
                zip_path = Path(__file__).resolve().parents[1] / zip_path
            if not zip_path.is_file():
                raise ValueError(f"{label} 图组尚未生成完成。")
            with zipfile.ZipFile(zip_path) as source:
                for name in source.namelist():
                    archive.writestr(f"{safe_name(label)}/{Path(name).name}", source.read(name))
            count += 1
    filename = f"{safe_name(group['product_name'])}_全部图组.zip"
    return filename, output.getvalue(), count


def build_output_archive(batch_id, ledger_path=LEDGER_PATH, store_id=STORE_ID):
    ensure_ledger(ledger_path)
    return _build_output_archive(batch_id, ledger_path, store_id=store_id)


def package_shared_data(ledger_path=LEDGER_PATH, files_dir=FILES_DIR, source_sheets_dir=SOURCE_SHEETS_DIR):
    ensure_ledger(ledger_path)
    result = _package_shared_data(ledger_path, files_dir, source_sheets_dir)
    data = load_ledger(ledger_path)
    missing = set(result.get("missing", []))
    for batch in data.get("batches", []):
        for row in batch.get("rows", []):
            for raw in list(row.get("image_reference_files", []) or []) + [row.get("image_zip", "")]:
                if raw and not Path(raw).is_absolute():
                    path = Path(__file__).resolve().parents[1] / raw
                else:
                    path = Path(raw).expanduser() if raw else None
                if path and not path.is_file():
                    missing.add(str(path))
            for group in (row.get("image_groups", {}) or {}).values():
                paths = list(group.get("reference_files", []) or [])
                paths.extend(group.get("generated_files", []) or [])
                paths.append(group.get("zip", ""))
                for raw in paths:
                    if not raw:
                        continue
                    path = Path(raw)
                    if not path.is_absolute():
                        path = Path(__file__).resolve().parents[1] / path
                    if not path.is_file():
                        missing.add(str(path))
    result["missing"] = sorted(missing)
    return result


def _normalize_workbench(data, ledger_path):
    batches = data.get("batches", [])
    if not batches:
        return ensure_ledger(ledger_path)

    batch = batches[0]
    changed = len(batches) != 1
    for field, value in {"id": BATCH_ID, "name": BATCH_NAME, "stores": [STORE_ID]}.items():
        if batch.get(field) != value:
            batch[field] = value
            changed = True
    metadata = _source_row_metadata(batch.get("source_sheet"))
    unique_rows = []
    seen_product_ids = set()
    for row in batch.get("rows", []):
        row_number = row.get("row_number")
        if row_number in RETAIN_DUPLICATE_SOURCE_ROWS and row.get("id") != f"ROW-{row_number}":
            row["id"] = f"ROW-{row_number}"
            changed = True
        product_id = (
            f"source-row-{row_number}"
            if row_number in RETAIN_DUPLICATE_SOURCE_ROWS
            else str(row.get("asin") or row.get("id") or "").strip().casefold()
        )
        if product_id and product_id in seen_product_ids:
            changed = True
            continue
        if product_id:
            seen_product_ids.add(product_id)
        if row.get("store_id") != STORE_ID:
            row["store_id"] = STORE_ID
            changed = True
        if row.get("purchase_status"):
            row["purchase_status"] = ""
            changed = True
        row_metadata = metadata.get(row.get("row_number"), {})
        for key in ("source_msku", "variant_size"):
            value = row_metadata.get(key, "")
            if value and row.get(key) != value:
                row[key] = value
                changed = True
        unique_rows.append(row)
    batch["rows"] = unique_rows
    data["batches"] = [batch]
    if changed:
        save_ledger(data, ledger_path)
    return data


def _source_row_metadata(source_sheet):
    path = Path(str(source_sheet or ""))
    if not path.is_absolute():
        path = Path(__file__).resolve().parents[1] / path
    if not path.is_file():
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = {
        str(sheet.cell(1, column).value or "").strip().casefold(): column
        for column in range(1, sheet.max_column + 1)
    }
    msku_column = headers.get("msku")
    metadata = {}
    if msku_column:
        for row_number in range(2, sheet.max_row + 1):
            msku = str(sheet.cell(row_number, msku_column).value or "").strip()
            size = msku.rsplit("-", 1)[-1].upper() if "-" in msku else ""
            metadata[row_number] = {
                "source_msku": msku,
                "variant_size": size if size in SIZE_ORDER else "",
            }
    workbook.close()
    return metadata


def _merge_confirmed_rows(rows):
    merged = []
    grouped = {}
    for row in rows:
        group = MERGE_GROUP_BY_ROW.get(row.get("row_number"))
        if group is None:
            merged.append(row)
            continue
        group_id = group["id"]
        value = group.get("values", {}).get(row.get("row_number"), "")
        primary = grouped.get(group_id)
        link_info = {
            "url": row.get("link", ""),
            "label": value,
            "row_number": row.get("row_number"),
            "row_id": row.get("id", ""),
        }
        template_info = {
            "file": row.get("files", {}).get("source_template", {}),
            "label": value,
            "row_number": row.get("row_number"),
            "row_id": row.get("id", ""),
        }
        if primary is None:
            primary = row
            primary["product_name"] = group["product_name"]
            primary["merge_group"] = group_id
            primary["merge_type"] = group.get("merge_type", "variation")
            primary["variation_theme"] = group.get("variation_theme", "")
            primary["attribute_label"] = group["attribute_label"]
            primary["source_links"] = [link_info]
            primary["source_template_files"] = [template_info] if template_info["file"].get("exists") else []
            primary["row_numbers"] = [row.get("row_number")]
            primary["merged_row_ids"] = [row.get("id", "")]
            primary["variant_values"] = [value] if value else []
            if group.get("variation_theme") == "SIZE":
                primary["variant_sizes"] = [value] if value else []
            grouped[group_id] = primary
            merged.append(primary)
            continue
        primary["source_links"].append(link_info)
        if template_info["file"].get("exists"):
            primary["source_template_files"].append(template_info)
        primary["row_numbers"].append(row.get("row_number"))
        primary["merged_row_ids"].append(row.get("id", ""))
        if value:
            primary["variant_values"].append(value)
            if group.get("variation_theme") == "SIZE":
                primary["variant_sizes"].append(value)
    for group_id, row in grouped.items():
        group = next(item for item in MERGE_GROUPS if item["id"] == group_id)
        order = {row_number: index for index, row_number in enumerate(group["rows"])}
        row["source_links"] = sorted(row["source_links"], key=lambda item: order.get(item.get("row_number"), 99))
        row["source_template_files"] = sorted(
            row["source_template_files"], key=lambda item: order.get(item.get("row_number"), 99)
        )
        row["row_numbers"] = sorted(set(row["row_numbers"]))
        row["variant_values"] = [
            group["values"][row_number]
            for row_number in group["rows"]
            if row_number in row["row_numbers"] and group["values"].get(row_number)
        ]
        if row.get("variation_theme") == "SIZE":
            row["variant_sizes"] = sorted(set(row["variant_sizes"]), key=lambda size: SIZE_ORDER.get(size, 99))
        if row.get("merge_type") == "bundle":
            row["image_group_specs"] = [{
                "key": "bundle",
                "label": "组合装",
                "source_url": "",
                "source_row_id": row.get("id", ""),
                "use_legacy": True,
            }]
        else:
            row["image_group_specs"] = [
                {
                    "key": item.get("row_id", ""),
                    "label": item.get("label", "") or f"子体{index}",
                    "source_url": item.get("url", ""),
                    "source_row_id": item.get("row_id", ""),
                    "use_legacy": item.get("row_id") == row.get("id"),
                }
                for index, item in enumerate(row["source_links"], start=1)
            ]
    return merged


def _resolve_image_group(batch_id, row_id, image_group_key, ledger_path):
    data = load_ledger(ledger_path)
    batch = next((item for item in data.get("batches", []) if item.get("id") == batch_id), None)
    if batch is None:
        raise ValueError("找不到批次。")
    rows = batch.get("rows", [])
    requested = next((item for item in rows if item.get("id") == row_id), None)
    if requested is None:
        raise ValueError("找不到产品行。")
    merge_group = MERGE_GROUP_BY_ROW.get(requested.get("row_number"))
    if merge_group is None:
        key = str(image_group_key or "default")
        if key != "default":
            raise ValueError("这个商品只有一套图组。")
        product_name = requested.get("product_name") or requested.get("asin") or row_id
        return row_id, "default", "", product_name

    members = [item for item in rows if item.get("row_number") in merge_group["rows"]]
    if not members:
        raise ValueError("合并商品没有可用的源行。")
    primary = members[0]
    if requested.get("id") != primary.get("id"):
        raise ValueError("请从合并商品主行操作图组。")
    if merge_group.get("merge_type", "variation") == "bundle":
        key = str(image_group_key or "bundle")
        if key not in {"default", "bundle"}:
            raise ValueError("组合装只使用一套组合图组。")
        return primary["id"], "bundle", "组合装", merge_group["product_name"]

    key = str(image_group_key or "")
    member = next((item for item in members if item.get("id") == key), None)
    if member is None:
        raise ValueError("找不到对应的子体图组。")
    label = merge_group.get("values", {}).get(member.get("row_number"), "")
    return primary["id"], key, label, merge_group["product_name"]
