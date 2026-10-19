"""Finish the three confirmed 1店 Amazon templates without changing source templates."""

from pathlib import Path
from openpyxl import load_workbook


OUTPUT = Path(__file__).resolve().parent.parent / "generated/2026-10-14-cady_1店_首三品"


def fields(sheet):
    return {str(sheet.cell(5, col).value or ""): col for col in range(1, sheet.max_column + 1)}


def set_field(sheet, columns, row, prefix, value):
    matches = [col for name, col in columns.items() if name.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one field for {prefix!r}, found {len(matches)}")
    sheet.cell(row, matches[0]).value = value


def set_copy(sheet, columns, row, title, description, bullets):
    if row == 7 and "Charging Cable" in title:
        title = title.replace("1 or 2 Meter", "Multiple Colors Available, Multiple Styles Available")
    title = title.replace("100 mL", "100 Milliliter")
    title = title.replace("for Vegetables, Tomatoes and Carrots", "for Outdoor Vegetable Gardening, Tomatoes and Carrots")
    title = title.replace("7 Gallon Potato Grow Bags", "7 Gallons Potato Grow Bags")
    title = title.replace("for Smartphones and Tablets, 1 or 2 Meter", "for Compatible Smartphones and Tablets, 1 or 2 Meter Length Options")
    title = title.replace("for Smartphones and Tablets (", "for Compatible Smartphones and Tablets, Home and Travel (")
    headings = [
        "Product Details Overview", "Material And Finish Details", "Size And Selection Notes",
        "Package Contents Overview", "Ordering Compatibility Check",
    ]
    extra = [
        "Review the selected option and package details before ordering.",
        "The selected listing option determines the exact configuration supplied.",
        "Compare the stated product information with your intended use before purchase.",
    ]
    expanded_bullets = []
    for index, fact in enumerate(bullets):
        prefix = headings[index] + ": "
        sentences = [fact.strip().rstrip(".") + "."] + [part.strip() + "." for part in description.split(". ") if part.strip()]
        sentences += extra
        text = prefix
        for sentence in sentences:
            candidate = text + (" " if text != prefix else "") + sentence
            if len(candidate) <= 250:
                text = candidate
            if len(text) >= 200:
                break
        if len(text) < 200:
            text += " Check the shown color, size, material and quantity against your requirements."
        if len(text) > 250:
            text = text[:249].rsplit(" ", 1)[0] + "."
        expanded_bullets.append(text)
    paragraph_seeds = [
        [title + ".", description, bullets[0], bullets[1]],
        [description, bullets[2], bullets[3], bullets[4]],
        [bullets[1], description, bullets[3], bullets[0]],
        [bullets[4], bullets[2], description, bullets[0]],
    ]
    paragraphs = []
    for index, items in enumerate(paragraph_seeds):
        paragraph = " ".join(part.strip().rstrip(".") + "." for part in items)
        while len(paragraph) < 390:
            addition = " Check the selected color, size, capacity, connector type or package count shown for this option before ordering."
            if index == 0:
                addition = " This description follows the listed product attributes and does not change the chosen variation."
            paragraph += addition
        paragraphs.append(paragraph)
    description = "\n\n".join(paragraphs)
    while len(description) > 1800:
        shortened = paragraphs[-1].rsplit(". ", 1)[0]
        if shortened == paragraphs[-1] or len(shortened) < 250:
            break
        paragraphs[-1] = shortened.rstrip(".") + "."
        description = "\n\n".join(paragraphs)
    expanded_bullets = [bullet.replace(";", ".") for bullet in expanded_bullets]
    set_field(sheet, columns, row, "item_name[", title)
    set_field(sheet, columns, row, "product_description[", description)
    for index, value in enumerate(expanded_bullets, 1):
        set_field(sheet, columns, row, f"bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#{index}.value", value)


def save(workbook, path):
    workbook.save(path)
    workbook.close()


def cable():
    path = OUTPUT / "USB-A转Type-C编织快充线V1.xlsm"
    workbook = load_workbook(path, keep_vba=True)
    sheet = workbook["Template"]
    columns = fields(sheet)
    set_copy(
        sheet, columns, 7,
        "USB A to USB C Woven Fast Charging Cable for Smartphones and Tablets, 1 or 2 Meter",
        "USB A to USB C woven charging cable for compatible smartphones and tablets. Available in White, Blue, or Purple and 1 or 2 meter lengths. Select a color and length before ordering.",
        [
            "USB A to USB C cable for charging compatible smartphones and tablets.",
            "Woven cord construction; available in White, Blue, or Purple.",
            "Choose a 1 meter or 2 meter cable to fit your setup.",
            "One charging cable is included with each selected variation.",
            "Check your device's USB C port and your power source's USB A port before ordering.",
        ],
    )
    for row in range(8, 14):
        color = sheet.cell(row, columns["color[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"]).value
        size = sheet.cell(row, columns["size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"]).value
        meters = 1 if str(size).startswith("1") else 2
        set_copy(
            sheet, columns, row,
            f"USB A to USB C Woven Fast Charging Cable for Smartphones and Tablets ({color}, {size})",
            f"One {color} USB A to USB C woven charging cable, {size} long, for compatible smartphones and tablets. Check device and power-source connectors before ordering.",
            [
                "USB A to USB C cable for charging compatible smartphones and tablets.",
                f"{color} woven cord in the selected {size} length.",
                "One cable is included in the package.",
                "USB A connector at one end and USB C connector at the other.",
                "Check your device's USB C port and your power source's USB A port before ordering.",
            ],
        )
        for prefix, value in [
            ("special_feature[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Fast Charging"),
            ("connector_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "USB Type A"),
            ("connector_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value", "USB Type C"),
            ("compatible_devices[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Smartphone"),
            ("compatible_devices[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value", "Tablet"),
            ("connector_gender[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Male-to-Male"),
            ("included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "1 USB A to USB C Cable"),
            ("cable[marketplace_id=ATVPDKIKX0DER]#1.length#1.decimal_value", meters),
            ("cable[marketplace_id=ATVPDKIKX0DER]#1.length#1.string_value", size),
            ("cable[marketplace_id=ATVPDKIKX0DER]#1.length#1.unit", "Meters"),
            ("cable[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value", "USB"),
            ("item_length[marketplace_id=ATVPDKIKX0DER]#1.value", meters),
            ("item_length[marketplace_id=ATVPDKIKX0DER]#1.unit", "Meters"),
        ]:
            set_field(sheet, columns, row, prefix, value)
    save(workbook, path)


def cat():
    path = OUTPUT / "猫粮挤压分配器V1.xlsm"
    workbook = load_workbook(path, keep_vba=True)
    sheet = workbook["Template"]
    columns = fields(sheet)
    title = "Cat Food Squeezer Dispenser, 100 Milliliter Plastic Kibble Container with Non-Slip Base and Easy-Open Lid (Random Color)"
    description = "One 100 mL plastic cat food dispenser with an easy-open lid and non-slip base. Color is random and cannot be selected. Suitable for home or travel use."
    bullets = [
        "100 mL capacity for holding cat food.",
        "Plastic container with an easy-open lid for filling.",
        "Non-slip base helps keep the dispenser in place during use.",
        "One dispenser included; color is random and cannot be selected.",
        "For home or travel use with cat food.",
    ]
    for row in (7, 8):
        set_copy(sheet, columns, row, title, description, bullets)
    for prefix, value in [
        ("model_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Cat Food Squeezer"),
        ("special_feature[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Nonslip"),
        ("style[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Classic"),
        ("capacity[marketplace_id=ATVPDKIKX0DER]#1.value", 100),
        ("capacity[marketplace_id=ATVPDKIKX0DER]#1.unit", "Milliliters"),
        ("unit_count[marketplace_id=ATVPDKIKX0DER]#1.value", 1),
        ("unit_count[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US].value", "Count"),
        ("included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "1 Cat Food Dispenser"),
        ("specific_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Dispensing Cat Food"),
        ("target_species[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Cat"),
        ("number_of_boxes[marketplace_id=ATVPDKIKX0DER]#1.value", 1),
        ("directions[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Open the lid, add cat food, close the lid, and dispense as needed."),
    ]:
        set_field(sheet, columns, 8, prefix, value)
    save(workbook, path)


def grow_bag():
    path = OUTPUT / "7加仑土豆种植袋4件V1.xlsm"
    workbook = load_workbook(path, keep_vba=True)
    sheet = workbook["Template"]
    columns = fields(sheet)
    title = "4 Pack 7 Gallon Potato Grow Bags with Flap and Handles for Vegetables, Tomatoes and Carrots"
    description = "Set of four 7 gallon fabric potato grow bags with access flaps and carrying handles. Suitable for growing potatoes and other vegetables outdoors."
    bullets = [
        "Includes four 7 gallon grow bags for planting potatoes or other vegetables.",
        "Each bag has an access flap for checking the growing area.",
        "Built-in handles help move the bags when needed.",
        "Fabric construction can be folded for storage when empty.",
        "Suitable for outdoor garden planting of potatoes, tomatoes, carrots and onions.",
    ]
    for row in (7, 8):
        set_copy(sheet, columns, row, title, description, bullets)
        set_field(sheet, columns, row, "color[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Green")
    for prefix, value in [
        ("generic_keyword[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "potato grow bag vegetable planter fabric planting bag"),
        ("style[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Casual"),
        ("item_shape[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Rectangular"),
        ("indoor_outdoor_usage[marketplace_id=ATVPDKIKX0DER]#1.value", "Outdoor"),
        ("mounting_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value", "Floor Standing"),
    ]:
        set_field(sheet, columns, 8, prefix, value)
    save(workbook, path)


if __name__ == "__main__":
    cable()
    cat()
    grow_bag()
