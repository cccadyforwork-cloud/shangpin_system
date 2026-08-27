from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/Users/cc/Desktop/8月w4上品/换土垫增加变体/换土垫多色多尺寸V1.xlsx")
OUTPUT = Path("/Users/cc/Documents/GitHub/shangpin_system/outputs/换土垫多色多尺寸V2.xlsx")

PACKAGE_DEFAULTS = {
    "item_package_dimensions[marketplace_id=ATVPDKIKX0DER]#1.length.value": 6,
    "item_package_dimensions[marketplace_id=ATVPDKIKX0DER]#1.width.value": 4,
    "item_package_dimensions[marketplace_id=ATVPDKIKX0DER]#1.height.value": 0.5,
    "item_package_dimensions[marketplace_id=ATVPDKIKX0DER]#1.length.unit": "Inches",
    "item_package_dimensions[marketplace_id=ATVPDKIKX0DER]#1.width.unit": "Inches",
    "item_package_dimensions[marketplace_id=ATVPDKIKX0DER]#1.height.unit": "Inches",
    "item_package_weight[marketplace_id=ATVPDKIKX0DER]#1.value": 0.25,
    "item_package_weight[marketplace_id=ATVPDKIKX0DER]#1.unit": "Pounds",
    "item_display_weight[marketplace_id=ATVPDKIKX0DER]#1.value": 0.25,
    "item_display_weight[marketplace_id=ATVPDKIKX0DER]#1.unit": "Pounds",
}

INSPECT_FIELDS = [
    "contribution_sku#1.value",
    "parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value",
    "color[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
    "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.depth.value",
    "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.height.value",
    "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.width.value",
    "item_package_dimensions[marketplace_id=ATVPDKIKX0DER]#1.length.value",
    "item_package_dimensions[marketplace_id=ATVPDKIKX0DER]#1.width.value",
    "item_package_dimensions[marketplace_id=ATVPDKIKX0DER]#1.height.value",
    "item_package_weight[marketplace_id=ATVPDKIKX0DER]#1.value",
]


def field_columns(ws):
    return {
        str(ws.cell(5, col).value).strip(): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(5, col).value
    }


def read_rows(ws, cols):
    sku_col = cols["contribution_sku#1.value"]
    rows = []
    for row in range(7, ws.max_row + 1):
        sku = ws.cell(row, sku_col).value
        if not sku:
            continue
        rows.append({
            field: ws.cell(row, cols[field]).value
            for field in INSPECT_FIELDS
            if field in cols
        })
    return rows


def main():
    wb = load_workbook(SOURCE)
    ws = wb["Template"]
    cols = field_columns(ws)

    print("Before:")
    for item in read_rows(ws, cols):
        print(item)

    parentage_col = cols["parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value"]
    sku_col = cols["contribution_sku#1.value"]
    for row in range(7, ws.max_row + 1):
        if not ws.cell(row, sku_col).value:
            continue
        if str(ws.cell(row, parentage_col).value or "").strip().lower() == "parent":
            continue
        for field, value in PACKAGE_DEFAULTS.items():
            if field in cols:
                ws.cell(row, cols[field]).value = value

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    wb2 = load_workbook(OUTPUT, data_only=True, read_only=True)
    ws2 = wb2["Template"]
    cols2 = field_columns(ws2)
    print("After:")
    for item in read_rows(ws2, cols2):
        print(item)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
