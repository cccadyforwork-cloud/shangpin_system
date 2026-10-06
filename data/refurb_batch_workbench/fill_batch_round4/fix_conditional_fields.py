from pathlib import Path

from openpyxl import load_workbook


OUTPUTS = Path(__file__).resolve().parent / "outputs"
SKU_FIELD = "contribution_sku#1.value"


def fill(file_name, values_by_sku):
    path = OUTPUTS / file_name
    wb = load_workbook(path, keep_vba=True)
    ws = wb["Template"]
    fields = {
        str(ws.cell(5, col).value).strip(): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(5, col).value
    }
    for row in range(7, ws.max_row + 1):
        sku = str(ws.cell(row, fields[SKU_FIELD]).value or "").strip()
        for field, value in values_by_sku.get(sku, {}).items():
            ws.cell(row, fields[field]).value = value
    wb.save(path)


fill(
    "冰袖4件装V1.xlsm",
    {
        "TTCA-CoolingSleeve": {
            "fabric_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Polyester",
        },
        "TTCA-CoolingSleeve-4pcs": {
            "model_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Cooling Sleeve",
            "style[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Sport",
            "department[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Unisex",
            "target_gender[marketplace_id=ATVPDKIKX0DER]#1.value": "Unisex",
            "age_range_description[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Adult",
            "fabric_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Polyester",
            "special_size_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Standard",
            "size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "One Size",
            "care_instructions[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Machine Wash",
            "closure[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value": "Pull-On",
            "import_designation[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Imported",
        },
    },
)

fill(
    "调料瓶蓝色2件装V1.xlsm",
    {
        "TTCA-SpiceShaker": {
            "item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "2 Pack 4 Oz Glass Salt and Pepper Shakers, Stainless Steel Lid Refillable Spice Jars for Kitchen Table RV Camp",
        },
        "TTCA-SpiceShaker-Blue-2pcs": {
            "item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "2 Pack 4 Oz Glass Salt and Pepper Shakers, Stainless Steel Lid Refillable Spice Jars for Kitchen Table RV Camp, Blue",
            "model_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Spice Shaker",
            "size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "4 Ounces",
            "care_instructions[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Hand Wash",
            "closure_material[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Stainless Steel",
            "included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "2 Spice Shakers",
        },
    },
)

fill(
    "清洁果蔬白灰色4件装V1.xlsm",
    {
        "TTCA-CleaningBrush-WhiteGrey-4pcs": {
            "model_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Double-Sided Cleaning Brush",
            "size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "3.78 x 3.78 x 6.3 Inches",
            "care_instructions[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Hand Wash",
            "closure_material[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "Plastic",
            "included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": "4 Double-Sided Cleaning Brushes",
        },
    },
)

print("conditional fields filled")
